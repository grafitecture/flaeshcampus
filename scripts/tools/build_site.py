#!/usr/bin/env python3
"""Baut die Vercel-Seite (site/) aus Episode 2, der Layout-Galerie und den Download-Dateien.

Aufruf (nach make_template.py --gallery und build_deck.py):
    python3 scripts/tools/build_site.py
Erzeugt site/index.html, site/slides/…, site/downloads/… und die Schriftdateien. CSS und JS liegen fest in site/assets/.
"""
import html
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image

REPO = Path(__file__).resolve().parents[2]
SKILL = REPO / "skill" / "flaesh-campus"
sys.path.insert(0, str(SKILL / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import PP_PLACEHOLDER  # noqa: E402

from campus.render import render_pdf  # noqa: E402
from campus.tokens import PH_IDX  # noqa: E402

SITE = REPO / "site"
FULL_W, THUMB_W = 1600, 400
STAND = "20. Oktober 2026"


# ---------------------------------------------------------------- Inhalte aus dem Deck

def _placeholder_text(slide, idx: int | None = None, title: bool = False) -> str:
    for ph in slide.placeholders:
        is_title = ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        if (title and is_title) or (not title and ph.placeholder_format.idx == idx):
            if ph.has_text_frame:
                return " ".join(p.text for p in ph.text_frame.paragraphs if p.text).strip()
    return ""


def slide_records(pptx: Path) -> list[dict]:
    """Nummer, Titel, Kicker, aktives Kapitel (aus der Leiste) und Sprechertext je Folie."""
    records = []
    for n, slide in enumerate(Presentation(str(pptx)).slides, 1):
        title = (_placeholder_text(slide, title=True) or _placeholder_text(slide, PH_IDX["Quote"])
                 or _placeholder_text(slide, PH_IDX["Question"]) or _placeholder_text(slide, PH_IDX["Statement"]))
        active = [s.name for s in slide.shapes if s.name.startswith("rail-dot-") and s.name.endswith("-active")]
        chapter = int(active[0].split("-")[2]) if active else None
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        records.append({"n": n, "title": title, "kicker": _placeholder_text(slide, PH_IDX["Kicker"]),
                        "chapter": chapter, "notes": notes, "layout": slide.slide_layout.name})
    return records


# ---------------------------------------------------------------- Bilder

def render_images(pptx: Path, out_dir: Path) -> list[tuple[str, str]]:
    """Rendert alle Folien als JPEG (1600 px) und Vorschau (400 px). Gibt relative Pfade zurück."""
    (out_dir / "thumb").mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        pdf = render_pdf(pptx, Path(tmp))
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(FULL_W), "-scale-to-y", "-1", str(pdf),
                        str(Path(tmp) / "p")], check=True)
        pages = sorted(Path(tmp).glob("p-*.png"))
        paths = []
        for i, page in enumerate(pages, 1):
            im = Image.open(page).convert("RGB")
            full, thumb = out_dir / f"{i:02d}.jpg", out_dir / "thumb" / f"{i:02d}.jpg"
            im.save(full, "JPEG", quality=84, optimize=True, progressive=True)
            im.resize((THUMB_W, round(im.height * THUMB_W / im.width)), Image.LANCZOS).save(
                thumb, "JPEG", quality=80, optimize=True)
            paths.append((full.relative_to(SITE).as_posix(), thumb.relative_to(SITE).as_posix()))
    return paths


def build_fonts(out_dir: Path) -> None:
    """Carlito (OFL) als WOFF-Teilmenge: humanistische Grotesk, der Corporate S der Folien am nächsten."""
    from fontTools import subset

    out_dir.mkdir(parents=True, exist_ok=True)
    unicodes = "U+0020-007E,U+00A0-00FF,U+2010-2027,U+2030,U+2039,U+203A,U+20AC,U+2190-2193,U+2212,U+2248"
    for weight, src in (("400", "Carlito-Regular.ttf"), ("700", "Carlito-Bold.ttf")):
        subset.main([f"/usr/share/fonts/truetype/crosextra/{src}", f"--unicodes={unicodes}", "--flavor=woff",
                     "--layout-features=kern,liga,tnum,lnum", f"--output-file={out_dir / f'carlito-{weight}.woff'}"])
    (out_dir / "Carlito-OFL.txt").write_text(
        "Carlito – Copyright (c) 2010-2013 by tyPoland Lukasz Dziedzic with Reserved Font Name \"Carlito\".\n"
        "Licensed under the SIL Open Font License, Version 1.1 – https://openfontlicense.org\n", encoding="utf-8")


# ---------------------------------------------------------------- HTML

def _size(path: Path) -> str:
    return f"{path.stat().st_size / 1_000_000:.1f} MB".replace(".", ",")


def _logo() -> str:
    svg = (SKILL / "assets" / "brand" / "logo.svg").read_text(encoding="utf-8")
    svg = svg[svg.index("<svg"):]
    return svg.replace("<svg", '<svg class="logo" role="img" aria-label="fläsh"', 1)


def write_index(out_dir: Path, episode: dict, slides: list[dict], layouts: list[dict], downloads: list[dict]) -> Path:
    e = html.escape
    chapters = "".join(
        f'<li><button type="button" class="rail-chapter" data-chapter="{i}" aria-label="Kapitel {i}: {e(name)}" '
        f'title="{e(name)}"><span class="dot" aria-hidden="true"></span><span class="num">{i}</span></button></li>'
        for i, name in enumerate(episode["chapters"], 1))
    thumbs = "".join(
        f'<li><button type="button" class="thumb" data-index="{i}" aria-label="Folie {s["n"]}: {e(s["title"])}">'
        f'<img src="{s["thumb"]}" alt="" width="400" height="225" loading="lazy" decoding="async"></button></li>'
        for i, s in enumerate(slides))
    layout_items = "".join(
        f'<li><button type="button" class="layout" data-src="{lay["src"]}" data-label="{e(lay["name"])}">'
        f'<img src="{lay["thumb"]}" alt="Layout {e(lay["name"])}" width="400" height="225" loading="lazy" decoding="async">'
        f'<span class="layout-name">{e(lay["name"])}</span><span class="layout-key">{e(lay["key"])}</span></button></li>'
        for lay in layouts)
    files = "".join(
        f'<li><a class="file" href="{d["href"]}" download><span class="file-label">{e(d["label"])}</span>'
        f'<span class="file-meta">{d["kind"]}, {d["size"]}</span><span class="file-action">Herunterladen</span></a></li>'
        for d in downloads)
    first = slides[0]
    data = json.dumps({"chapters": episode["chapters"], "slides": slides}, ensure_ascii=False).replace("</", "<\\/")
    page = f"""<!doctype html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex, nofollow">
<title>fläsh Campus · PPT</title>
<meta name="description" content="PPT zur Trainingsreihe fläsh Campus: {e(episode["title"])} mit {e(episode["presenter"])}.">
<link rel="preload" href="assets/fonts/carlito-400.woff" as="font" type="font/woff" crossorigin>
<link rel="stylesheet" href="assets/site.css">
<link rel="icon" href="assets/favicon.svg" type="image/svg+xml">
</head>
<body>
<a class="skip" href="#viewer">Zur Präsentation springen</a>
<aside class="rail" aria-label="Kapitel der PPT">
  <p class="rail-ep">PPT</p>
  <ol class="rail-chapters">{chapters}</ol>
  {_logo()}
</aside>
<main>
  <section class="stage" id="folien" aria-labelledby="titel">
    <header class="masthead">
      <p class="brand">{_logo()}<span class="campus">Campus</span></p>
      <nav class="sections" aria-label="Bereiche"><a href="#folien">Folien</a><a href="#layouts">Layouts</a><a href="#dateien">Dateien</a></nav>
    </header>
    <div class="intro">
      <h1 id="titel">{e(episode["title"])}</h1>
      <p>Die PPT zur Trainingsreihe mit {e(episode["presenter"])}. Blättern mit den Pfeiltasten oder über die Vorschau.</p>
    </div>
    <figure class="viewer" id="viewer" tabindex="-1">
      <div class="screen"><img id="slide" src="{first["src"]}" alt="Folie 1: {e(first["title"])}" width="1600" height="900" fetchpriority="high"></div>
      <figcaption class="controls">
        <button type="button" class="step" id="prev" aria-label="Vorherige Folie"><svg viewBox="0 0 24 24" aria-hidden="true"><path d="M15 5l-7 7 7 7"/></svg></button>
        <span class="counter" aria-live="polite"><span id="current">1</span> / {len(slides)}</span>
        <button type="button" class="step" id="next" aria-label="Nächste Folie"><svg viewBox="0 0 24 24" aria-hidden="true"><path d="M9 5l7 7-7 7"/></svg></button>
        <span class="where" id="where"></span>
        <button type="button" class="notes-toggle" id="notes-toggle" aria-expanded="false" aria-controls="notes">Sprechertext anzeigen</button>
      </figcaption>
      <div class="notes" id="notes" hidden><p id="notes-text"></p></div>
    </figure>
    <ol class="thumbs" aria-label="Alle Folien">{thumbs}</ol>
  </section>
  <section class="paper" id="layouts" aria-labelledby="layouts-titel">
    <h2 id="layouts-titel">Die 18 Layouts</h2>
    <p class="lead">Aus diesen Layouts entsteht jede Folie der Reihe. Die Beispiele zeigen Inhalte aus der PPT. Der Schlüssel unter dem Namen hilft bei Korrekturen, etwa „Folie 7 als compare“.</p>
    <ul class="layouts">{layout_items}</ul>
  </section>
  <section class="paper files-section" id="dateien" aria-labelledby="dateien-titel">
    <h2 id="dateien-titel">Dateien</h2>
    <ul class="files">{files}</ul>
    <p class="hint">Zum Bearbeiten der PowerPoint-Dateien muss die Schrift Corporate S installiert sein.</p>
  </section>
  <footer class="colophon"><p>fläsh Campus ist eine Trainingsreihe von WHITEsmile mit {e(episode["presenter"])}. Gestaltung und Umsetzung: Grafitecture. Stand: {STAND}.</p></footer>
</main>
<dialog class="lightbox" id="lightbox" aria-labelledby="lightbox-label">
  <form method="dialog"><button class="close" aria-label="Schließen">×</button></form>
  <img id="lightbox-img" src="" alt="">
  <p id="lightbox-label"></p>
</dialog>
<script type="application/json" id="deck-data">{data}</script>
<script src="assets/site.js" defer></script>
</body>
</html>
"""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "index.html"
    path.write_text(page, encoding="utf-8")
    return path


# ---------------------------------------------------------------- Ablauf

def _layout_records(gallery: Path, paths: list[tuple[str, str]]) -> list[dict]:
    seen, layouts = set(), []
    for rec, (src, thumb) in zip(slide_records(gallery), paths):
        name = rec["layout"]
        key = rec["kicker"].split("·")[-1].strip() if "·" in rec["kicker"] else ""
        if name in seen:
            name += " mit Fotos" if "photos" in key else " (Auflösung)"
            key = key.replace(" with photos", " + photo")
        seen.add(rec["layout"])
        layouts.append({"name": name, "key": key, "src": src, "thumb": thumb})
    return layouts


def download_sources(deck: Path, gallery: Path) -> list[tuple[str, Path, str, str]]:
    """Je Download: Beschriftung, Quelldatei, Dateiname auf der Seite, Dateiart."""
    return [
        ("PPT als PDF", REPO / "out" / "episode-02.pdf", "flaesh-campus-in-office-whitening-workflow.pdf", "PDF"),
        ("PPT als PowerPoint", deck, "flaesh-campus-in-office-whitening-workflow.pptx", "PowerPoint"),
        ("Vorlage mit allen Layouts", gallery, "flaesh-campus-template.pptx", "PowerPoint"),
        ("Layout-Übersicht als PDF", REPO / "out" / "flaesh-campus-layouts.pdf", "flaesh-campus-layouts.pdf", "PDF"),
    ]


def main() -> None:
    plan = json.loads((SKILL / "examples" / "episode-02.json").read_text(encoding="utf-8"))
    ep = plan["episode"]
    deck, gallery = REPO / "out" / "episode-02.pptx", SKILL / "assets" / "template" / "flaesh-campus-template.pptx"
    for d in ("slides", "downloads"):
        shutil.rmtree(SITE / d, ignore_errors=True)
    records = slide_records(deck)
    for rec, (src, thumb) in zip(records, render_images(deck, SITE / "slides" / "ep02")):
        rec.update(src=src, thumb=thumb)
        rec.pop("layout")
    layouts = _layout_records(gallery, render_images(gallery, SITE / "slides" / "layouts"))
    dl = SITE / "downloads"
    dl.mkdir(parents=True)
    downloads = []
    for label, src, name, kind in download_sources(deck, gallery):
        shutil.copyfile(src, dl / name)
        downloads.append({"label": label, "href": f"downloads/{name}", "size": _size(dl / name), "kind": kind})
    build_fonts(SITE / "assets" / "fonts")
    episode = {"number": ep["number"], "title": ep["title"], "presenter": ep["presenter"]["name"],
               "chapters": ep["chapters"]}
    print(write_index(SITE, episode, records, layouts, downloads))


if __name__ == "__main__":
    main()
