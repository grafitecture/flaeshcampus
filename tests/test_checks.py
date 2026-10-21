import copy
import json
import shutil
import subprocess

import pytest
from pptx import Presentation
from pptx.util import Pt

from campus.build import build_deck
from campus.checks import check_deck
from campus.measure import Measurer
from campus.render import render_pdf
from campus.template import layout_spec
from campus.tokens import PH_IDX, px
from conftest import REPO, SKILL

FIXTURE = REPO / "tests" / "fixtures" / "minimal-plan.json"


@pytest.fixture(scope="module")
def m(font_path):
    return Measurer(font_path)


@pytest.fixture
def plan():
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def build(plan, template_path, tmp_path, name="deck.pptx"):
    out = tmp_path / name
    build_deck(plan, template_path, out, SKILL / "examples")
    return out


def kinds(issues, kind):
    return [i for i in issues if i.kind == kind]


def test_clean_minimal(plan, template_path, tmp_path, m):
    issues = check_deck(build(plan, template_path, tmp_path), layout_spec(), m)
    assert issues == []


def test_zone_violation(plan, template_path, tmp_path, m):
    path = build(plan, template_path, tmp_path)
    prs = Presentation(path)
    tb = prs.slides[2].shapes.add_textbox(px(1600), px(800), px(200), px(60))
    tb.name = "intruder"
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "x"
    run.font.name = "Corporate S"
    prs.save(path)
    zone = kinds(check_deck(path, layout_spec(), m), "zone")
    assert [(i.slide, i.shape) for i in zone] == [(3, "intruder")]


def test_video_zone_is_not_an_intruder(plan, template_path, tmp_path, m):
    """Der Video-Kreis über der Vollbild-Fotofolie gehört in die Zone und ist kein Fehler."""
    plan["slides"].insert(2, {"type": "photo", "kicker": "Step by step", "title": "Prepare & Protect",
                              "photo": {"file": "photos/ep02/07-adjust-lamp.jpg"}, "caption": "Lamp docked."})
    path = build(plan, template_path, tmp_path)
    assert [s.name for s in Presentation(path).slides[2].shapes][-1] == "zone-presenter-top"
    assert kinds(check_deck(path, layout_spec(), m), "zone") == []


def test_overflow(plan, template_path, tmp_path, m):
    long_text = ("Light intensity falls with the square of distance so one centimetre of extra gap between "
                 "lamp and teeth costs roughly half of the available light power at the enamel surface")
    plan["slides"][2]["points"] = [{"lead": f"Point {i}", "text": long_text} for i in range(4)]
    issues = check_deck(build(plan, template_path, tmp_path), layout_spec(), m)
    assert any(i.kind == "overflow" and i.slide == 3 and i.shape == "Points" for i in issues)


def test_glyph(plan, template_path, tmp_path, m):
    plan["slides"][2]["title"] = "Do → Don't ✓"
    issues = check_deck(build(plan, template_path, tmp_path), layout_spec(), m)
    glyphs = kinds(issues, "glyph")
    assert glyphs and "→" in glyphs[0].detail and "✓" in glyphs[0].detail


def test_font(plan, template_path, tmp_path, m):
    path = build(plan, template_path, tmp_path)
    prs = Presentation(path)
    tb = prs.slides[2].shapes.add_textbox(px(240), px(880), px(400), px(40))
    run = tb.text_frame.paragraphs[0].add_run()
    run.text = "Arial text"
    run.font.name = "Arial"
    run.font.size = Pt(13)
    prs.save(path)
    assert kinds(check_deck(path, layout_spec(), m), "font")


WIDGET_SLIDES = [
    {"type": "points_chart", "title": "The Lamp: Position Is Power",
     "points": [{"lead": "Lock the geometry.", "text": "Patient at 45°, lamp locked on the retractor."},
                {"lead": "The expensive centimetre.", "text": "One centimetre of gap costs ≈50% of the power."}],
     "chart": {"title": "Light power at the enamel",
               "bars": [{"label": "Lamp docked on the retractor", "value": 100, "highlight": True},
                        {"label": "1 cm of extra gap", "value": 56}],
               "note": "1 cm of extra gap ≈ 50% less light power"}},
    {"type": "scale", "title": "The Aftercare Window",
     "timeline": {"points": [{"at": "Minutes", "label": "Basal layer forms"},
                             {"at": "~2 h", "label": "Near-full thickness"},
                             {"at": "18–24 h", "label": "Maturation in vivo"}]}},
    {"type": "summary", "title": "The SOP Blueprint", "items": [
        {"title": "Protect", "text": "Vitamin E or vaseline, correct retractor, goggles, cured barrier — avoid the Seven Sins."},
        {"title": "Activate & monitor", "text": "Lamp locked on the retractor · fläshcard discipline · comfort checkpoints."},
        {"title": "Seal & instruct", "text": "Mousse, the aftercare window, the departure gift — identically, every time."}]},
    {"type": "closing", "title": "Make It Your Team’s Reflex", "steps": [
        {"title": "Run one full dry run", "text": "No patient in the chair — your assistant drives the fläshcard start to finish."},
        {"title": "Appoint an SOP owner", "text": "One team member owns this protocol — and audits the clinic against it."},
        {"title": "Laminate the phases", "text": "Visible chairside, next to the Seven Sins of Improper Isolation."}],
     "next": {"title": "Episode 03 · Sensitivity", "text": "Preventing it, managing it."}},
]


def test_widgets_fit_their_boxes(plan, template_path, tmp_path, m):
    """Einzeilige Beschriftungen, dreizeilige Zusammenfassungen und zweizeilige Schritte passen in ihre Boxen."""
    plan["slides"] = copy.deepcopy(WIDGET_SLIDES)
    path = build(plan, template_path, tmp_path)
    assert kinds(check_deck(path, layout_spec(), m), "overflow") == []
    shapes = {s.name: s for s in Presentation(path).slides[0].shapes}
    for i in (1, 2):
        label, track = shapes[f"bar-label-{i}"], shapes[f"bar-track-{i}"]
        assert label.top + label.height <= track.top


def test_timeline_caption_below_labels(plan, template_path, tmp_path, m):
    """Zweizeilige Zeitstrahl-Beschriftungen und die Bildunterschrift überlappen nicht."""
    plan["slides"] = [{"type": "scale", "title": "The Aftercare Window", "caption": "The pellicle rebuilds fast",
                       "timeline": {"points": [{"at": "Minutes", "label": "Basal pellicle layer forms quickly"},
                                               {"at": "~2 h", "label": "Near-full thickness"},
                                               {"at": "18–24 h", "label": "Maturation in vivo"}]},
                       "points": [{"lead": "Mind the first hours.", "text": "Chromogen-light choices."}]}]
    path = build(plan, template_path, tmp_path)
    assert check_deck(path, layout_spec(), m) == []
    shapes = {s.name: s for s in Presentation(path).slides[0].shapes}
    caption = shapes["scale-caption"]
    assert all(caption.top >= s.top + s.height for n, s in shapes.items() if n.startswith("tl-label-"))
    points = next(s for s in Presentation(path).slides[0].placeholders if s.placeholder_format.idx == PH_IDX["Points"])
    assert caption.top + caption.height <= points.top


def test_photo_steps_slide_clean(plan, template_path, tmp_path, m):
    """Eine Schrittfolie mit Fotos und Bildunterschrift besteht die Prüfung; ein fehlendes Foto wird Briefing."""
    from PIL import Image

    for i in range(2):
        Image.new("RGB", (1920, 1080), "white").save(tmp_path / f"step{i}.jpg")
    plan["slides"] = [{"type": "steps", "title": "Seal & Finish", "start": 9,
                       "caption": "After Whitening Mousse soothes and remineralises — then final shade & photos.",
                       "steps": [{"title": "Rinse off residual gel with water", "photo": {"file": "step0.jpg"}},
                                 {"title": "Remove the Gingiva Protector", "photo": {"file": "step1.jpg"}},
                                 {"title": "Apply After Whitening Mousse", "photo": {"brief": "Mousse applied"}}]}]
    out = tmp_path / "steps.pptx"
    report = build_deck(plan, template_path, out, tmp_path)
    assert check_deck(out, layout_spec(), m) == []
    assert report.briefs == [(1, "Mousse applied")]


def test_chapter_title_two_lines(plan, template_path, tmp_path, m):
    plan["slides"] = [{"type": "chapter", "chapter": n, "line": "The last ten minutes decide how the first sixty are remembered."}
                      for n in (4, 5)]
    path = build(plan, template_path, tmp_path)
    assert check_deck(path, layout_spec(), m) == []
    for slide in Presentation(path).slides:
        title = next(s for s in slide.placeholders if s.placeholder_format.idx == 0)
        line = next(s for s in slide.placeholders if s.placeholder_format.idx == PH_IDX["Line"])
        assert line.top >= title.top + title.height
        assert line.top + line.height <= px(1000)


def test_long_chapter_words_reported(plan, template_path, tmp_path, m):
    """Überlange Wörter im Kapiteltitel werden gemeldet, und nichts rutscht unter die Folie."""
    plan["episode"]["chapters"] = ["Document", "Patient Communication & Aftercare", "Contraindications"]
    plan["slides"] = [{"type": "chapter", "chapter": n, "line": "One sentence for the chapter."} for n in (2, 3)]
    out = tmp_path / "long.pptx"
    report = build_deck(plan, template_path, out, SKILL / "examples")
    issues = check_deck(out, layout_spec(), m)
    words = kinds(issues, "word")
    assert any(i.slide == 1 and "Communication" in i.detail for i in words)
    assert any(i.slide == 2 and "Contraindications" in i.detail for i in words)
    assert kinds(issues, "bounds") == []
    assert len(report.warnings) == 2
    for slide in Presentation(out).slides:
        for shape in slide.shapes:
            assert shape.top + shape.height <= px(1080), shape.name


def test_bounds_reported(plan, template_path, tmp_path, m):
    path = build(plan, template_path, tmp_path)
    prs = Presentation(path)
    tb = prs.slides[2].shapes.add_textbox(px(240), px(1060), px(400), px(60))
    tb.name = "below"
    tb.text_frame.paragraphs[0].add_run().text = "x"
    prs.save(path)
    bounds = kinds(check_deck(path, layout_spec(), m), "bounds")
    assert [(i.slide, i.shape) for i in bounds] == [(3, "below")]


@pytest.mark.slow
def test_render_pages(plan, template_path, tmp_path):
    if not shutil.which("pdfinfo"):
        pytest.skip("pdfinfo missing")
    path = build(plan, template_path, tmp_path)
    pdf = render_pdf(path, tmp_path / "render")
    info = subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
    pages = int(next(line.split()[-1] for line in info.splitlines() if line.startswith("Pages:")))
    assert pages == len(Presentation(path).slides)


def test_check_cli(plan, template_path, tmp_path):
    import sys

    path = build(plan, template_path, tmp_path)
    cli = [sys.executable, str(SKILL / "scripts" / "check_deck.py"), str(path)]
    ok = subprocess.run(cli, capture_output=True, text=True)
    assert ok.returncode == 0, ok.stdout + ok.stderr
    prs = Presentation(path)
    tb = prs.slides[2].shapes.add_textbox(px(1600), px(800), px(200), px(60))
    tb.name = "intruder"
    tb.text_frame.paragraphs[0].add_run().text = "x"
    prs.save(path)
    bad = subprocess.run(cli, capture_output=True, text=True)
    assert bad.returncode == 1
    assert "slide 3 zone intruder" in bad.stdout
