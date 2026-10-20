import pathlib
import subprocess

import pytest
from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

from campus.tokens import SLIDE_EMU, ZONE, px

NAMES = [
    "01 Title", "02 Chapter", "03 Hook", "04 Closing", "05 Points", "06 Points + Photo",
    "07 Points + Chart", "08 Statement", "09 Checklist", "10 Steps", "11 Compare", "12 Scale",
    "13 Summary", "14 Photo", "15 Before / After", "16 Photo Annotated", "17 Case", "18 Quiz",
]
KICKER_LAYOUTS = NAMES[2:13] + NAMES[14:18]          # 03–13, 15–18
TITLE_LAYOUTS = NAMES[3:7] + NAMES[8:13] + NAMES[14:17]  # 04–07, 09–13, 15–17 (08 hat Big statt Titel)
VALIDATE = pathlib.Path("/mnt/skills/public/pptx/scripts/office/validate.py")


@pytest.fixture(scope="module")
def prs(template_path):
    return Presentation(template_path)


def layout(prs, name):
    return next(lay for lay in prs.slide_layouts if lay.name == name)


def named(lay, name):
    return next((s for s in lay.shapes if s.name == name), None)


def box(shape):
    return (shape.left, shape.top, shape.width, shape.height)


def test_slide_size(prs):
    assert (prs.slide_width, prs.slide_height) == SLIDE_EMU


def test_layout_names(prs):
    assert [lay.name for lay in prs.slide_layouts] == NAMES


def test_theme_fonts(prs):
    theme = prs.slide_master.part.part_related_by(RT.THEME)
    xml = etree.fromstring(theme.blob)
    major = xml.find(".//" + qn("a:majorFont")).find(qn("a:latin")).get("typeface")
    minor = xml.find(".//" + qn("a:minorFont")).find(qn("a:latin")).get("typeface")
    assert major == minor == "Corporate S"


def test_zone_free_without_ring(prs):
    """Kein sichtbarer Kreis unten rechts (Wunsch Martin, 19.10.); die Fläche bleibt für das Video reserviert."""
    from campus.template import layout_spec

    spec = layout_spec()
    for name in NAMES:
        lay = layout(prs, name)
        assert named(lay, "zone-ring") is None, name
        assert not [s.name for s in lay.shapes if box(s) == tuple(px(v) for v in ZONE)], name
        assert spec[name]["zone"] is (name != "01 Title"), name


def test_rail(prs):
    for name in NAMES:
        rail = named(layout(prs, name), "rail")
        assert rail is not None, name
        assert rail.left == 0 and rail.width == px(150), name


def test_header(prs):
    for name in KICKER_LAYOUTS:
        lay = layout(prs, name)
        kicker = next(p for p in lay.placeholders if p.placeholder_format.idx == 11)
        assert (kicker.left, kicker.top) == (px(240), px(110)), name
    for name in TITLE_LAYOUTS:
        lay = layout(prs, name)
        title = next(p for p in lay.placeholders if p.placeholder_format.type == PP_PLACEHOLDER.TITLE)
        assert box(title) == (px(240), px(150), px(1500), px(100)), name
    for name in ("03 Hook", "08 Statement", "18 Quiz"):
        lay = layout(prs, name)
        assert not any(p.placeholder_format.type == PP_PLACEHOLDER.TITLE for p in lay.placeholders), name


def test_validate(template_path):
    if not VALIDATE.exists():
        pytest.skip("pptx validator not available")
    result = subprocess.run(["python3", str(VALIDATE), str(template_path)], capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr
