import json

import pytest
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from campus.template import layout_spec, write_layout_spec
from campus.tokens import PH_IDX, SIZE_PT, px

TYPES = {"title": PP_PLACEHOLDER.TITLE, "body": PP_PLACEHOLDER.BODY, "pic": PP_PLACEHOLDER.PICTURE}

EXPECTED = [
    ("01 Title", "Kicker", "body", (230, 356, 980, 40), "kicker"),
    ("01 Title", "Title", "title", (226, 404, 960, 230), "display"),
    ("01 Title", "Subtitle", "body", (230, 654, 820, 110), "subtitle"),
    ("01 Title", "Presenter", "body", (230, 880, 1000, 46), "lead"),
    ("01 Title", "Credentials", "body", (230, 930, 1000, 80), "small"),
    ("01 Title", "Photo", "pic", (1236, 250, 580, 580), "small"),
    ("02 Chapter", "Number", "body", (230, 20, 560, 880), "chapter_number"),
    ("02 Chapter", "Kicker", "body", (800, 330, 900, 40), "kicker"),
    ("02 Chapter", "Title", "title", (792, 376, 1000, 180), "chapter_title"),
    ("02 Chapter", "Line", "body", (800, 590, 700, 130), "line"),
    ("03 Hook", "Quote", "body", (240, 170, 1300, 300), "quote"),
    ("03 Hook", "TruthLabel", "body", (290, 560, 500, 30), "label"),
    ("03 Hook", "TruthText", "body", (290, 610, 500, 240), "body"),
    ("03 Hook", "AnswerLabel", "body", (930, 560, 520, 30), "label"),
    ("03 Hook", "AnswerText", "body", (930, 610, 520, 240), "body"),
    ("04 Closing", "NextLabel", "body", (1230, 370, 520, 30), "label"),
    ("04 Closing", "NextTitle", "body", (1230, 410, 520, 120), "lead"),
    ("04 Closing", "NextText", "body", (1230, 540, 520, 110), "body"),
    ("05 Points", "Points", "body", (240, 330, 1260, 560), "lead"),
    ("06 Points + Photo", "Points", "body", (240, 330, 700, 560), "lead"),
    ("06 Points + Photo", "Photo", "pic", (1000, 300, 800, 400), "small"),
    ("06 Points + Photo", "Caption", "body", (1000, 712, 540, 40), "small"),
    ("07 Points + Chart", "Points", "body", (240, 330, 700, 560), "lead"),
    ("08 Statement", "Big", "body", (240, 200, 1300, 330), "statement_big"),
    ("08 Statement", "Statement", "body", (240, 560, 1200, 150), "statement"),
    ("08 Statement", "Text", "body", (240, 730, 1200, 130), "body"),
    ("12 Scale", "Points", "body", (240, 560, 1260, 330), "lead"),
    ("14 Photo", "Photo", "pic", (150, 0, 1770, 1080), "small"),
    ("14 Photo", "Kicker", "body", (190, 780, 1300, 60), "kicker"),
    ("14 Photo", "Title", "title", (190, 840, 1300, 110), "title"),
    ("14 Photo", "Caption", "body", (190, 950, 1300, 90), "body"),
    ("15 Before / After", "PhotoBefore", "pic", (240, 330, 600, 400), "small"),
    ("15 Before / After", "PhotoAfter", "pic", (900, 330, 600, 400), "small"),
    ("15 Before / After", "LabelBefore", "body", (240, 750, 600, 40), "body"),
    ("15 Before / After", "LabelAfter", "body", (900, 750, 600, 40), "body"),
    ("16 Photo Annotated", "Photo", "pic", (240, 300, 900, 600), "small"),
    ("17 Case", "Story", "body", (240, 330, 1260, 150), "lead"),
    ("18 Quiz", "Question", "body", (240, 180, 1260, 220), "quote"),
]


@pytest.fixture(scope="module")
def prs(template_path):
    return Presentation(template_path)


def find(prs, layout_name, name, ph_type):
    lay = next(lay for lay in prs.slide_layouts if lay.name == layout_name)
    for ph in lay.placeholders:
        if ph_type == "title":
            if ph.placeholder_format.type == PP_PLACEHOLDER.TITLE:
                return ph
        elif ph.placeholder_format.idx == PH_IDX[name]:
            return ph
    return None


@pytest.mark.parametrize("layout_name,name,ph_type,box,size_key", EXPECTED)
def test_placeholder_matrix(prs, layout_name, name, ph_type, box, size_key):
    ph = find(prs, layout_name, name, ph_type)
    assert ph is not None, f"{layout_name}: {name} missing"
    assert ph.placeholder_format.type == TYPES[ph_type]
    actual = (ph.left, ph.top, ph.width, ph.height)
    for a, e in zip(actual, (px(v) for v in box)):
        assert abs(a - e) <= 1, (layout_name, name, actual)
    d = ph._element.txBody.find(qn("a:lstStyle")).find(qn("a:lvl1pPr")).find(qn("a:defRPr"))
    assert d.get("sz") == str(int(SIZE_PT[size_key] * 100))


def test_cover_photo(prs):
    ph = find(prs, "01 Title", "Photo", "pic")
    sp_pr = ph._element.spPr
    assert sp_pr.find(qn("a:prstGeom")).get("prst") == "ellipse"
    assert sp_pr.find(qn("a:ln")).get("w") == str(px(6))
    assert sp_pr.find(qn("a:effectLst")).find(qn("a:glow")).get("rad") == str(px(90))


def test_spec_json(tmp_path):
    path = tmp_path / "layout-spec.json"
    write_layout_spec(path)
    spec = json.loads(path.read_text(encoding="utf-8"))
    assert len(spec) == 18
    assert spec == json.loads(json.dumps(layout_spec()))
    for layout_name, entry in spec.items():
        assert {"key", "frame", "zone", "placeholders"} <= set(entry), layout_name
        for name, ph in entry["placeholders"].items():
            assert {"idx", "type", "box", "levels"} <= set(ph), (layout_name, name)
