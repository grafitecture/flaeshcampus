import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from campus.tokens import COLOR, PH_IDX, px
from campus.widgets_content import (bars_card, case_findings, checklist, closing_steps, compare_columns, markers,
                                    points_text, quiz_options, steps_grid, summary_items, timeline)

SKY = RGBColor.from_string(COLOR["sky"])


@pytest.fixture
def prs(template_path):
    return Presentation(template_path)


def new_slide(prs, layout_name):
    return prs.slides.add_slide(next(lay for lay in prs.slide_layouts if lay.name == layout_name))


def by_name(slide):
    return {s.name: s for s in slide.shapes}


def all_text(slide):
    return " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


def test_points_runs(prs):
    slide = new_slide(prs, "05 Points")
    ph = slide.placeholders[PH_IDX["Points"]]
    points_text(ph, [{"lead": "Lock the geometry", "text": "Patient at 45°, lamp docked on the retractor."},
                     {"lead": "Light, not heat", "text": "Blue 460 nm light activates the green gel."}])
    paras = ph.text_frame.paragraphs
    assert paras[0].text == "01\tLock the geometry"
    assert paras[0].runs[0].font.color.rgb == SKY
    assert paras[0].space_before == Pt(0)
    assert paras[1].level == 1
    assert paras[2].text == "02\tLight, not heat"


def test_checklist_rows(prs):
    slide = new_slide(prs, "09 Checklist")
    checklist(slide, [f"Item {i}" for i in range(1, 8)])
    shapes = by_name(slide)
    tops = [shapes[f"check-num-{i}"].top for i in range(1, 8)]
    assert all(b - a == px(82) for a, b in zip(tops, tops[1:]))
    assert shapes["check-text-7"].text_frame.text == "Item 7"


@pytest.mark.parametrize("n", [4, 5])
def test_steps_grid(prs, n):
    slide = new_slide(prs, "10 Steps")
    steps_grid(slide, [{"title": f"Step {i}", "text": "Text"} for i in range(n)], dark=False)
    tops = {s.top for s in slide.shapes if s.name.startswith("step-num-")}
    assert len(tops) == 2
    assert len([s for s in slide.shapes if s.name.startswith("step-num-")]) == n


def test_bars(prs):
    slide = new_slide(prs, "07 Points + Chart")
    bars_card(slide, {"title": "Light power at the tooth",
                      "bars": [{"label": "Docked · 3 cm", "value": 100, "highlight": True},
                               {"label": "+1 cm gap · 4 cm", "value": 56}],
                      "note": "Intensity falls with distance²: 3² = 9 vs. 4² = 16"})
    shapes = by_name(slide)
    assert abs(shapes["bar-fill-2"].width - round(0.56 * shapes["bar-track-2"].width)) <= 1
    assert shapes["bar-fill-1"].fill.fore_color.rgb == SKY
    card = shapes["chart-card"]
    assert card.top + card.height <= px(724) or card.left + card.width <= px(1564)


def test_bar_labels_above_fills(prs):
    """Beschriftungen liegen in der Ebenenfolge über den Balken (Leuchten verdeckt keinen Text); Werte ohne Leerzeichen."""
    slide = new_slide(prs, "07 Points + Chart")
    bars_card(slide, {"title": "Light power", "bars": [{"label": "Docked", "value": 100, "highlight": True},
                                                       {"label": "+1 cm", "value": 56}]})
    order = [s.name for s in slide.shapes]
    for i in (1, 2):
        assert order.index(f"bar-label-{i}") > order.index(f"bar-fill-{i}")
    assert by_name(slide)["bar-value-2"].text_frame.text == "56%"


def glow_px(shape) -> float:
    from pptx.oxml.ns import qn

    return int(shape._element.spPr.find(qn("a:effectLst")).find(qn("a:glow")).get("rad")) / 6350


def test_timeline_labels_do_not_overlap(prs):
    slide = new_slide(prs, "12 Scale")
    timeline(slide, [{"at": f"{i} h", "label": "Near-full thickness"} for i in range(6)])
    shapes = by_name(slide)
    for i in range(1, 6):
        a, b = shapes[f"tl-label-{i}"], shapes[f"tl-label-{i + 1}"]
        assert a.left + a.width <= b.left
    dot = shapes["tl-dot-1"]
    assert glow_px(dot) <= dot.width / 6350 / 2


def test_marker_glow_small(prs):
    slide = new_slide(prs, "16 Photo Annotated")
    markers(slide, (240, 300, 900, 600), [{"x": 0.5, "y": 0.5, "label": "Retractor docked"}])
    dot = by_name(slide)["marker-1"]
    assert glow_px(dot) <= dot.width / 6350 / 2


def bottom(shape):
    return shape.top + shape.height


def test_summary_rows_breathe(prs):
    """Dreizeilige Texte stoßen nicht an den nächsten Punkt; die letzte Zeile endet über der Evidence."""
    slide = new_slide(prs, "13 Summary")
    summary_items(slide, [{"title": f"Item {i}", "text": "Text"} for i in range(5)])
    shapes = by_name(slide)
    assert shapes["sum-title-2"].top - bottom(shapes["sum-text-1"]) >= px(16)
    assert bottom(shapes["sum-text-3"]) <= px(920)


def test_closing_rows_breathe(prs):
    slide = new_slide(prs, "04 Closing")
    closing_steps(slide, [{"title": f"Step {i}", "text": "Text"} for i in range(3)])
    shapes = by_name(slide)
    assert shapes["next-step-num-2"].top - bottom(shapes["next-step-text-1"]) >= px(16)


def test_case_rows_clear_story_and_evidence(prs):
    slide = new_slide(prs, "17 Case")
    case_findings(slide, [{"finding": f"Finding {i}", "cause": "Cause"} for i in range(5)])
    shapes = by_name(slide)
    assert shapes["finding-head"].top >= px(330 + 150 + 24)
    assert bottom(shapes["cause-5"]) <= px(920)


def test_quiz_highlight(prs):
    slide = new_slide(prs, "18 Quiz")
    quiz_options(slide, ["Gel below 1 mm", "Barrier over-extended", "Lamp gap"], 1, "Explanation.")
    shapes = by_name(slide)
    assert shapes["quiz-opt-2"].fill.fore_color.rgb == SKY
    for i in (1, 3):
        assert shapes[f"quiz-opt-{i}"].fill.fore_color.rgb == RGBColor.from_string(COLOR["night2"])
    assert shapes["quiz-expl"].text_frame.text == "Explanation."


def test_no_glyphs(prs):
    slide = new_slide(prs, "11 Compare")
    compare_columns(slide, {"label": "Myth", "items": ["Whitening etches the enamel"]},
                    {"label": "Fact", "items": ["Oxidation of chromogens at neutral pH"]}, "myth_fact")
    slide2 = new_slide(prs, "13 Summary")
    summary_items(slide2, [{"title": "Document", "text": "Clean first."}, {"title": "Protect", "text": "Barrier."}])
    for s in (slide, slide2):
        text = all_text(s)
        assert "✓" not in text and "→" not in text
    assert any(sh.name.startswith("icon-check-r-") for sh in slide.shapes)
    assert any(sh.name.startswith("icon-cross-l-") for sh in slide.shapes)
