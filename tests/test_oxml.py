import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.shapes.shapetree import SlideShapes

from campus.oxml import add_glow, set_radial_fill, to_placeholder
from campus.tokens import px


@pytest.fixture
def layout_shapes():
    prs = Presentation()
    layout = prs.slide_layouts[6]
    return SlideShapes(layout.shapes._spTree, layout)


def test_glow_after_ln(layout_shapes):
    shape = layout_shapes.add_shape(MSO_SHAPE.OVAL, px(100), px(100), px(300), px(300))
    shape.line.color.rgb = RGBColor.from_string("00A5FF")
    add_glow(shape, "00A5FF", 46, 45)
    tags = [child.tag for child in shape._element.spPr]
    assert tags.index(qn("a:ln")) < tags.index(qn("a:effectLst"))
    glow = shape._element.spPr.find(qn("a:effectLst")).find(qn("a:glow"))
    assert glow.get("rad") == str(px(46))
    assert glow.find(qn("a:srgbClr")).find(qn("a:alpha")).get("val") == "45000"


def test_radial_fill(layout_shapes):
    shape = layout_shapes.add_shape(MSO_SHAPE.OVAL, 0, 0, px(600), px(600))
    set_radial_fill(shape, "00A5FF", [(0, 34), (50, 10), (100, 0)])
    grad = shape._element.spPr.find(qn("a:gradFill"))
    assert grad.find(qn("a:path")).get("path") == "circle"
    stops = grad.find(qn("a:gsLst")).findall(qn("a:gs"))
    alphas = [gs.find(qn("a:srgbClr")).find(qn("a:alpha")).get("val") for gs in stops]
    assert alphas == ["34000", "10000", "0"]


def test_body_placeholder(layout_shapes):
    tb = layout_shapes.add_textbox(px(240), px(330), px(1260), px(560))
    to_placeholder(
        tb, "body", 26, "Points",
        levels=[{"size_pt": 21, "color": "5D6C73"}, {"size_pt": 17, "color": "333333", "mar_l_px": 80}],
        prompt=["Point title", "Explanation"],
    )
    assert tb.is_placeholder
    assert tb.placeholder_format.idx == 26
    lst = tb._element.txBody.find(qn("a:lstStyle"))
    lvl1 = lst.find(qn("a:lvl1pPr"))
    d = lvl1.find(qn("a:defRPr"))
    assert d.get("sz") == "2100"
    assert d.find(qn("a:solidFill")).find(qn("a:srgbClr")).get("val") == "5D6C73"
    assert d.find(qn("a:latin")).get("typeface") == "Corporate S"
    assert lst.find(qn("a:lvl2pPr")).get("marL") == str(px(80))


def test_title_placeholder(layout_shapes):
    tb = layout_shapes.add_textbox(px(240), px(150), px(1500), px(100))
    to_placeholder(tb, "title", None, "Title", levels=[{"size_pt": 39, "color": "5D6C73"}], prompt=["Slide title"])
    ph = tb._element.nvSpPr.nvPr.find(qn("p:ph"))
    assert ph.get("type") == "title"
    assert "idx" not in ph.attrib


def test_pic_placeholder_ellipse(layout_shapes):
    shape = layout_shapes.add_shape(MSO_SHAPE.RECTANGLE, px(1236), px(250), px(580), px(580))
    to_placeholder(shape, "pic", 15, "Photo", levels=[{"size_pt": 13, "color": "7D8B91"}],
                   prompt=["Insert photo"], geom="ellipse")
    ph = shape._element.nvSpPr.nvPr.find(qn("p:ph"))
    assert ph.get("type") == "pic"
    assert shape._element.spPr.find(qn("a:prstGeom")).get("prst") == "ellipse"


def test_fill_text_frame(layout_shapes):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt

    from campus.oxml import fill_text_frame

    tb = layout_shapes.add_textbox(0, px(84), px(150), px(30))
    style = {"size_pt": 11, "color": "ABB7BE", "caps": True, "spc": 176, "align": "center"}
    fill_text_frame(tb.text_frame, [("EP", style)], anchor="b")
    para = tb.text_frame.paragraphs[0]
    run = para.runs[0]
    assert tb.text_frame.text == "EP"
    assert run.font.size == Pt(11) and run.font.name == "Corporate S"
    assert run.font.color.rgb == RGBColor.from_string("ABB7BE")
    assert run._r.rPr.get("cap") == "all" and run._r.rPr.get("spc") == "176"
    assert para.alignment == PP_ALIGN.CENTER
    body_pr = tb.text_frame._txBody.bodyPr
    assert body_pr.get("anchor") == "b" and body_pr.get("wrap") == "square"
    assert body_pr.find(qn("a:spAutoFit")) is None
    assert tb.text_frame.margin_left == 0 and tb.text_frame.margin_top == 0
