import pytest
from PIL import Image
from pptx import Presentation

from campus.measure import Measurer
from campus.tokens import PH_IDX, px
from campus.widgets_frame import evidence, focus_box, footer, place_photo, rail_progress

CHAPTERS = ["Document", "Protect", "Apply", "Activate & Monitor", "Seal & Instruct"]


@pytest.fixture
def prs(template_path):
    return Presentation(template_path)


def new_slide(prs, layout_name):
    layout = next(lay for lay in prs.slide_layouts if lay.name == layout_name)
    return prs.slides.add_slide(layout)


def names(slide):
    return [s.name for s in slide.shapes]


def test_rail_active(prs):
    slide = new_slide(prs, "05 Points")
    rail_progress(slide, CHAPTERS, 3, dark=False)
    dots = [n for n in names(slide) if n.startswith("rail-dot-")]
    assert len(dots) == 5
    assert [n for n in dots if n.endswith("-active")] == ["rail-dot-3-active"]
    assert "rail-line" in names(slide)
    assert {f"rail-num-{i}" for i in range(1, 6)} <= set(names(slide))


def glow_px(shape) -> float:
    from pptx.oxml.ns import qn

    glow = shape._element.spPr.find(qn("a:effectLst")).find(qn("a:glow"))
    return int(glow.get("rad")) / 6350


def test_rail_active_glow_clear_of_label(prs):
    """Das Leuchten des aktiven Punkts bleibt klein (kein Leucht-Klotz) und lässt die Kapitelnummer frei."""
    slide = new_slide(prs, "05 Points")
    rail_progress(slide, CHAPTERS, 2, dark=False)
    shapes = {s.name: s for s in slide.shapes}
    dot, num = shapes["rail-dot-2-active"], shapes["rail-num-2"]
    assert glow_px(dot) <= dot.width / 6350 / 2
    assert dot.left + dot.width + px(glow_px(dot)) <= num.left


@pytest.mark.parametrize("n", [2, 6])
def test_rail_bounds(prs, n):
    slide = new_slide(prs, "05 Points")
    rail_progress(slide, [f"Chapter {i}" for i in range(n)], 1, dark=False)
    dots = [s for s in slide.shapes if s.name.startswith("rail-dot-")]
    assert len(dots) == n
    for dot in dots:
        assert px(300) <= dot.top <= px(900), dot.name
        assert dot.top + dot.height <= px(900), dot.name


def test_photo_aspect(prs, tmp_path):
    Image.new("RGB", (400, 1200), "white").save(tmp_path / "tall.jpg")
    slide = new_slide(prs, "06 Points + Photo")
    warnings = []
    assert place_photo(slide, PH_IDX["Photo"], {"file": "tall.jpg"}, tmp_path, warnings) is None
    pic = slide.placeholders[PH_IDX["Photo"]]
    assert pic.crop_top + pic.crop_bottom > 0
    assert pic.crop_left == pic.crop_right == 0
    assert warnings == []


def test_photo_unreadable(prs, tmp_path):
    (tmp_path / "broken.jpg").write_bytes(b"not an image at all")
    slide = new_slide(prs, "06 Points + Photo")
    warnings = []
    brief = place_photo(slide, PH_IDX["Photo"], {"file": "broken.jpg", "brief": "Lamp docked"}, tmp_path, warnings)
    assert brief == "Lamp docked"
    assert "brief-card" in names(slide)
    assert any("broken.jpg" in w for w in warnings)


def test_photo_webp_with_jpg_suffix(prs, tmp_path):
    Image.new("RGB", (1600, 800), "white").save(tmp_path / "mislabeled.jpg", "WEBP")
    slide = new_slide(prs, "06 Points + Photo")
    warnings = []
    assert place_photo(slide, PH_IDX["Photo"], {"file": "mislabeled.jpg"}, tmp_path, warnings) is None
    assert warnings == []


def test_photo_exif_orientation(prs, tmp_path):
    exif = Image.Exif()
    exif[0x0112] = 6                       # Kamera hochkant: um 90° drehen
    Image.new("RGB", (1600, 800), "white").save(tmp_path / "rotated.jpg", exif=exif)
    slide = new_slide(prs, "06 Points + Photo")
    place_photo(slide, PH_IDX["Photo"], {"file": "rotated.jpg"}, tmp_path, [])
    assert slide.placeholders[PH_IDX["Photo"]].image.size == (800, 1600)


def test_photo_too_small(prs, tmp_path):
    """Spec §6: zu kleines Foto → Platzhalter mit Briefing und Eintrag in der Fotoliste."""
    Image.new("RGB", (40, 24), "white").save(tmp_path / "thumb.jpg")
    slide = new_slide(prs, "06 Points + Photo")
    warnings = []
    brief = place_photo(slide, PH_IDX["Photo"], {"file": "thumb.jpg"}, tmp_path, warnings)
    assert brief and "brief-card" in names(slide)
    assert any("thumb.jpg" in w for w in warnings)


def test_photo_downscaled(prs, tmp_path):
    """4K-Fotos werden auf höchstens das Doppelte des Rahmens verkleinert, sonst wird das Deck riesig."""
    Image.new("RGB", (3840, 2160), "white").save(tmp_path / "uhd.jpg")
    slide = new_slide(prs, "06 Points + Photo")  # Rahmen 800 × 400
    place_photo(slide, PH_IDX["Photo"], {"file": "uhd.jpg"}, tmp_path, [])
    w, h = slide.placeholders[PH_IDX["Photo"]].image.size
    assert min(w / 800, h / 400) <= 2.01


def test_photo_fit_keeps_whole_image(prs, tmp_path):
    """Produktbilder und Bildschirme: ganz zeigen statt beschneiden, mittig im Rahmen."""
    Image.new("RGBA", (600, 720), (0, 0, 0, 0)).save(tmp_path / "product.png")
    slide = new_slide(prs, "06 Points + Photo")  # Rahmen x 1000, 800 × 400
    place_photo(slide, PH_IDX["Photo"], {"file": "product.png", "fit": True}, tmp_path, [])
    pic = slide.placeholders[PH_IDX["Photo"]]
    assert (pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom) == (0, 0, 0, 0)
    assert abs(pic.height - px(400)) <= px(1)
    assert abs(pic.width - px(400 * 600 / 720)) <= px(1)
    assert abs(pic.left + pic.width / 2 - px(1400)) <= px(1)


def test_placed_photo_drops_placeholder_fill(prs, tmp_path):
    """Transparente Produktbilder zeigen den Folienhintergrund, nicht die graue Leerfläche des Platzhalters."""
    from pptx.oxml.ns import qn

    Image.new("RGBA", (600, 720), (0, 0, 0, 0)).save(tmp_path / "product.png")
    slide = new_slide(prs, "06 Points + Photo")
    place_photo(slide, PH_IDX["Photo"], {"file": "product.png", "fit": True}, tmp_path, [])
    sp_pr = slide.placeholders[PH_IDX["Photo"]]._element.spPr
    assert sp_pr.find(qn("a:noFill")) is not None


def test_photo_missing(prs, tmp_path):
    slide = new_slide(prs, "06 Points + Photo")
    warnings = []
    brief = place_photo(slide, PH_IDX["Photo"], {"file": "nope.jpg", "brief": "Foto: Lampe"}, tmp_path, warnings)
    assert brief == "Foto: Lampe"
    assert "brief-card" in names(slide)
    assert any("nope.jpg" in w for w in warnings)
    assert PH_IDX["Photo"] not in [ph.placeholder_format.idx for ph in slide.placeholders]


def test_webp(prs, tmp_path):
    Image.new("RGB", (800, 400), "gray").save(tmp_path / "wide.webp", "WEBP")
    slide = new_slide(prs, "06 Points + Photo")
    place_photo(slide, PH_IDX["Photo"], {"file": "wide.webp"}, tmp_path, [])
    assert slide.placeholders[PH_IDX["Photo"]].image.content_type == "image/jpeg"


def test_focus_height(prs, font_path):
    m = Measurer(font_path)
    short = focus_box(new_slide(prs, "05 Points"), (1220, 330, 580), "Clinical focus",
                      "Burns are never caused by the gel.", "Short.", m)
    long = focus_box(new_slide(prs, "05 Points"), (1220, 330, 580), "Clinical focus",
                     "Burns are never caused by the gel.",
                     "Rescue: vitamin E or vaseline onto the gingiva — no panic, protocol. " * 3, m)
    assert long > short


def test_evidence_and_footer(prs):
    slide = new_slide(prs, "05 Points")
    evidence(slide, "Ontiveros & Paravina, J Dent 2009", dark=False)
    footer(slide, "IN-OFFICE WHITENING WORKFLOW · 12", dark=False)
    shapes = {s.name: s for s in slide.shapes}
    assert shapes["evidence-label"].text_frame.text == "Evidence"
    assert shapes["evidence-text"].text_frame.text == "Ontiveros & Paravina, J Dent 2009"
    assert (shapes["evidence-text"].left, shapes["evidence-text"].top) == (px(240), px(948))
    assert shapes["footer"].top == px(1010)


def test_brief_card_stays_below_text(prs, tmp_path):
    slide = new_slide(prs, "14 Photo")
    place_photo(slide, PH_IDX["Photo"], {"brief": "Foto: Isolation complete"}, tmp_path, [])
    order = [s.name for s in slide.shapes]
    title = next(s.name for s in slide.placeholders if s.placeholder_format.idx == 0)
    assert order.index("brief-card") < order.index(title)
    assert order.index("brief-text") < order.index(title)
