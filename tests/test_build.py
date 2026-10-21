import copy
import json
import subprocess
import sys

import pytest
from pptx import Presentation
from pptx.shapes.placeholder import PlaceholderPicture

from campus.build import build_deck
from campus.tokens import PH_IDX, px
from conftest import REPO, SKILL

EPISODE = {
    "number": 2,
    "title": "In-Office Whitening Workflow",
    "subtitle": "Step-by-Step Protocol + Common Mistakes",
    "presenter": {"name": "Dr. Germán Gómez", "credentials": "DDS · MD · PhD",
                  "photo": {"file": "photos/gomez.jpg"}},
    "chapters": ["Document", "Protect", "Apply", "Activate & Monitor", "Seal & Instruct"],
}
POINTS = [{"lead": "Lock the geometry", "text": "Patient at 45°, lamp docked on the retractor."},
          {"lead": "Light, not heat", "text": "Blue 460 nm light activates the green gel."}]
ALL_TYPES = {"episode": EPISODE, "slides": [
    {"type": "cover"},
    {"type": "chapter", "chapter": 3, "line": "The gel is engineered for precision."},
    {"type": "hook", "quote": "Will mine look like hers?",
     "truth": {"text": "Same gel, same lamp, different results."},
     "answer": {"text": "A standardized protocol ends the lottery."}},
    {"type": "closing", "title": "Make it your team's reflex",
     "steps": [{"title": "Run one dry run", "text": "No patient in the chair."},
               {"title": "Appoint an SOP owner", "text": "One person audits the protocol."}],
     "next": {"title": "Episode 03 · Sensitivity", "text": "Preventing and managing it."}},
    {"type": "points", "chapter": 3, "title": "The lamp: position is power", "points": POINTS,
     "evidence": "Ontiveros & Paravina, J Dent 2009", "notes": "Notes for slide five."},
    {"type": "points_photo", "chapter": 1, "title": "Clean before you whiten", "points": POINTS,
     "photo": {"brief": "Foto: Professional prophylaxis"}, "caption": "Before and after"},
    {"type": "points_chart", "chapter": 3, "title": "The expensive centimetre", "points": POINTS,
     "chart": {"title": "Light power at the tooth",
               "bars": [{"label": "Docked · 3 cm", "value": 100, "highlight": True},
                        {"label": "+1 cm gap · 4 cm", "value": 56}]}},
    {"type": "statement", "big": "30", "statement": "years of German-made whitening expertise"},
    {"type": "checklist", "chapter": 2, "title": "The seven sins of isolation",
     "items": ["No lip protection", "Wrong retractor", "No eye protection", "Barrier not cured",
               "Gel below 1 mm", "Patient unattended", "Sensitivity not managed"],
     "focus": {"title": "Burns are never caused by the gel.", "text": "Rescue: vitamin E onto the gingiva."}},
    {"type": "steps", "title": "The fix", "steps": [{"title": f"Step {i}", "text": "Do this."} for i in range(4)]},
    {"type": "compare", "title": "Oxidation, not erosion", "style": "myth_fact",
     "left": {"label": "Myth", "items": ["Whitening etches the enamel"]},
     "right": {"label": "Fact", "items": ["Oxidation of chromogens at neutral pH"]}},
    {"type": "scale", "chapter": 1, "title": "The value-ordered baseline",
     "scale": {"items": ["B1", "A1", "B2", "D2"], "highlight": ["A1"]}, "caption": "lightest to darkest",
     "points": POINTS},
    {"type": "summary", "title": "The SOP blueprint",
     "items": [{"title": "Document", "text": "Clean first."}, {"title": "Protect", "text": "Cured barrier."}]},
    {"type": "photo", "kicker": "Step by step", "title": "Prepare & Protect",
     "photo": {"brief": "Foto: Isolation complete"}, "caption": "The barrier is the treatment."},
    {"type": "before_after", "title": "Identical setup",
     "before": {"label": "Before · A3", "photo": {"brief": "Foto: baseline"}},
     "after": {"label": "After · B1", "photo": {"brief": "Foto: final shade"}}},
    {"type": "photo_annotated", "title": "Isolation setup", "photo": {"brief": "Foto: isolation"},
     "markers": [{"x": 0.3, "y": 0.4, "label": "Lip protection"}, {"x": 0.6, "y": 0.5, "label": "Retractor"}]},
    {"type": "case", "title": "The uneven case", "story": "A 38-year-old at the end of cycle 3.",
     "findings": [{"finding": "Cervical thirds still yellow", "cause": "Barrier over-extended"}]},
    {"type": "quiz", "question": "Most likely cause?", "options": ["Gel", "Barrier", "Lamp"], "answer": 1,
     "explanation": "The barrier covered cervical enamel."},
]}
EXPECTED_LAYOUTS = ["01 Title", "02 Chapter", "03 Hook", "04 Closing", "05 Points", "06 Points + Photo",
                    "07 Points + Chart", "08 Statement", "09 Checklist", "10 Steps", "11 Compare",
                    "12 Scale", "13 Summary", "14 Photo", "15 Before / After", "16 Photo Annotated",
                    "17 Case", "18 Quiz", "18 Quiz"]


@pytest.fixture(scope="module")
def deck(template_path, tmp_path_factory):
    out = tmp_path_factory.mktemp("deck") / "all.pptx"
    report = build_deck(copy.deepcopy(ALL_TYPES), template_path, out, SKILL / "examples")
    return Presentation(out), report


def test_layout_per_type(deck):
    prs, report = deck
    assert [s.slide_layout.name for s in prs.slides] == EXPECTED_LAYOUTS
    assert report.slides == 19


def test_optional_absent(deck, template_path, tmp_path):
    prs, _ = deck
    for n, slide in enumerate(prs.slides, 1):
        for ph in slide.placeholders:
            if isinstance(ph, PlaceholderPicture):
                continue
            assert ph.has_text_frame and ph.text_frame.text.strip(), (n, ph.name)
    plan = {"episode": EPISODE, "slides": [
        {"type": "points", "chapter": 3, "title": "With chapter", "points": POINTS},
        {"type": "points", "title": "Without chapter", "points": POINTS},
    ]}
    out = tmp_path / "optional.pptx"
    build_deck(plan, template_path, out, SKILL / "examples")
    slides = list(Presentation(out).slides)
    kicker = next(ph for ph in slides[0].placeholders if ph.placeholder_format.idx == PH_IDX["Kicker"])
    assert kicker.text_frame.text == "Phase 3 · Apply"
    assert PH_IDX["Kicker"] not in [ph.placeholder_format.idx for ph in slides[1].placeholders]


def test_video_zone_on_photo_slides(deck):
    """Nur die Vollbild-Fotofolie legt den Video-Kreis selbst obendrauf; alle anderen erben ihn vom Layout."""
    prs, _ = deck
    photo_slides = 0
    for n, slide in enumerate(prs.slides, 1):
        names = [s.name for s in slide.shapes]
        if slide.slide_layout.name == "14 Photo":
            photo_slides += 1
            assert names[-2:] == ["zone-ring-top", "zone-presenter-top"], n
        else:
            assert not [name for name in names if name.startswith("zone-")], n
    assert photo_slides


def test_notes(deck):
    prs, _ = deck
    assert prs.slides[4].notes_slide.notes_text_frame.text == "Notes for slide five."


def test_footer(deck):
    prs, _ = deck
    footer = next(s for s in prs.slides[4].shapes if s.name == "footer")
    assert footer.text_frame.text.endswith("· 5")


def test_briefs_reported(deck):
    _, report = deck
    assert (6, "Foto: Professional prophylaxis") in report.briefs


def test_cli_plan_error(tmp_path):
    plan = copy.deepcopy(ALL_TYPES)
    plan["slides"].append({"type": "bullets"})
    path = tmp_path / "plan.json"
    path.write_text(json.dumps(plan), encoding="utf-8")
    result = subprocess.run([sys.executable, str(SKILL / "scripts" / "build_deck.py"), str(path),
                             "--out", str(tmp_path / "x.pptx")], capture_output=True, text=True)
    assert result.returncode == 2
    assert "unknown type" in result.stdout


def test_make_template_gallery(tmp_path):
    """--gallery baut die Layout-Galerie in die Vorlage; Decks aus dieser Vorlage enthalten nur ihre eigenen Folien."""
    out = tmp_path / "template" / "flaesh-campus-template.pptx"
    gallery = SKILL / "examples" / "layout-gallery.json"
    result = subprocess.run([sys.executable, str(SKILL / "scripts" / "make_template.py"), "--out", str(out),
                             "--gallery", str(gallery)], capture_output=True, text=True)
    assert result.returncode == 0, result.stdout + result.stderr
    prs = Presentation(out)
    plan = json.loads(gallery.read_text(encoding="utf-8"))
    assert len(prs.slide_layouts) == 18
    assert len(prs.slides) == len(plan["slides"]) + 1  # das Quiz bringt seine Auflösungsfolie mit
    assert sorted({s.slide_layout.name for s in prs.slides}) == sorted(lay.name for lay in prs.slide_layouts)
    assert (out.parent / "layout-spec.json").exists()
    deck = tmp_path / "deck.pptx"
    build_deck({"episode": plan["episode"], "slides": [{"type": "cover"}]}, out, deck, SKILL / "examples")
    assert len(Presentation(deck).slides) == 1


def test_cli_missing_template(tmp_path):
    """Ein falscher --template-Pfad bricht ab, statt still eine frische Vorlage zu nehmen."""
    path = tmp_path / "plan.json"
    path.write_text(json.dumps({"episode": EPISODE, "slides": [{"type": "cover"}]}), encoding="utf-8")
    result = subprocess.run([sys.executable, str(SKILL / "scripts" / "build_deck.py"), str(path),
                             "--out", str(tmp_path / "x.pptx"), "--template", str(tmp_path / "nope.pptx")],
                            capture_output=True, text=True)
    assert result.returncode == 2
    assert "nope.pptx" in result.stdout + result.stderr
    assert not (tmp_path / "x.pptx").exists()


def test_markers_follow_moved_photo(template_path, tmp_path):
    """Verschiebt Martin das Foto in Layout 16, landen die Marker trotzdem auf dem Foto."""
    prs = Presentation(template_path)
    lay = next(lay for lay in prs.slide_layouts if lay.name == "16 Photo Annotated")
    ph = next(p for p in lay.placeholders if p.placeholder_format.idx == PH_IDX["Photo"])
    ph.left, ph.top, ph.width, ph.height = px(300), px(320), px(800), px(560)
    moved = tmp_path / "moved.pptx"
    prs.save(moved)
    plan = {"episode": EPISODE, "slides": [{"type": "photo_annotated", "title": "Isolation setup",
                                            "photo": {"brief": "Isolation setup"},
                                            "markers": [{"x": 0.5, "y": 0.5, "label": "Centre"}]}]}
    out = tmp_path / "deck.pptx"
    build_deck(plan, moved, out, SKILL / "examples")
    marker = next(s for s in Presentation(out).slides[0].shapes if s.name == "marker-1")
    assert abs(marker.left + marker.width // 2 - px(700)) <= px(1)
    assert abs(marker.top + marker.height // 2 - px(600)) <= px(1)


def test_cli_ok(tmp_path, template_path):
    plan = {"episode": {**EPISODE, "presenter": {"name": "Dr. Germán Gómez"}},
            "slides": [{"type": "cover"}, {"type": "points", "title": "Two points", "points": POINTS}]}
    path = tmp_path / "plan.json"
    path.write_text(json.dumps(plan), encoding="utf-8")
    result = subprocess.run([sys.executable, str(SKILL / "scripts" / "build_deck.py"), str(path),
                             "--out", str(tmp_path / "ok.pptx"), "--template", str(template_path)],
                            capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["slides"] == 2
