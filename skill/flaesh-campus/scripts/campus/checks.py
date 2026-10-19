"""Prüfungen für gebaute Decks: Überlauf, Kreis-Zone, leere Platzhalter, fremde Schriften, Glyphen."""
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from fontTools.ttLib import TTFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.shapes.placeholder import PlaceholderPicture

from .measure import Measurer, Para
from .tokens import EMU_PER_PX, FONT, PH_IDX, zone_circle

IDX_NAME = {idx: name for name, idx in PH_IDX.items()}
ALLOWED_FONTS = {None, FONT, "+mj-lt", "+mn-lt"}
ZONE_EXEMPT_ON_PHOTO = ("brief-",)


@dataclass
class Issue:
    slide: int
    kind: Literal["overflow", "zone", "empty", "font", "glyph"]
    shape: str
    detail: str


def _px(emu) -> float:
    return (emu or 0) / EMU_PER_PX


def _logical_name(shape) -> str:
    if shape.is_placeholder:
        return IDX_NAME.get(shape.placeholder_format.idx, shape.name)
    return shape.name


def _layout_ph(slide, idx: int):
    return next((p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == idx), None)


def _level_props(layout_ph, level: int) -> dict:
    """Effektive Absatz- und Zeichenwerte einer Ebene aus dem lstStyle des Layout-Platzhalters."""
    props = {"size_pt": None, "ln_pct": 100, "before": 0.0, "after": 0.0, "mar_l": 0.0, "spc": 0.0}
    if layout_ph is None:
        return props
    lst = layout_ph._element.txBody.find(qn("a:lstStyle")) if layout_ph._element.txBody is not None else None
    if lst is None:
        return props
    levels = [lst.find(qn(f"a:lvl{i}pPr")) for i in range(1, 10)]
    levels = [lv for lv in levels if lv is not None]
    if not levels:
        return props
    ppr = levels[min(level, len(levels) - 1)]
    props["mar_l"] = _px(int(ppr.get("marL", 0)))
    ln = ppr.find(qn("a:lnSpc"))
    if ln is not None and ln.find(qn("a:spcPct")) is not None:
        props["ln_pct"] = int(ln.find(qn("a:spcPct")).get("val")) / 1000
    for key, tag in (("before", "a:spcBef"), ("after", "a:spcAft")):
        el = ppr.find(qn(tag))
        if el is not None and el.find(qn("a:spcPts")) is not None:
            props[key] = int(el.find(qn("a:spcPts")).get("val")) / 100
    d = ppr.find(qn("a:defRPr"))
    if d is not None:
        if d.get("sz"):
            props["size_pt"] = int(d.get("sz")) / 100
        if d.get("spc"):
            props["spc"] = int(d.get("spc")) / 100
    return props


def _insets(shape, layout_ph) -> tuple[float, float, float, float]:
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    source = body
    if layout_ph is not None and not any(body.get(k) for k in ("lIns", "tIns", "rIns", "bIns")):
        source = layout_ph._element.txBody.find(qn("a:bodyPr"))
    defaults = {"lIns": 91440, "tIns": 45720, "rIns": 91440, "bIns": 45720}
    return tuple(_px(int(source.get(k, defaults[k]))) for k in ("lIns", "tIns", "rIns", "bIns"))


def _paras(shape, layout_ph) -> list[Para]:
    paras = []
    for p in shape.text_frame.paragraphs:
        props = _level_props(layout_ph, p.level)
        size = next((r.font.size.pt for r in p.runs if r.font.size is not None), None) or props["size_pt"] or 18
        spc = props["spc"]
        for r in p.runs:
            if r._r.rPr is not None and r._r.rPr.get("spc"):
                spc = int(r._r.rPr.get("spc")) / 100
        ln_pct = props["ln_pct"]
        if isinstance(p.line_spacing, float):
            ln_pct = p.line_spacing * 100
        before = p.space_before.pt if p.space_before is not None else props["before"]
        after = p.space_after.pt if p.space_after is not None else props["after"]
        paras.append(Para(p.text, size, ln_pct, before, after, spc, props["mar_l"]))
    return paras


def _intersects_zone(shape) -> bool:
    cx, cy, r = zone_circle()
    x, y, w, h = _px(shape.left), _px(shape.top), _px(shape.width), _px(shape.height)
    nx, ny = max(x, min(cx, x + w)), max(y, min(cy, y + h))
    return (nx - cx) ** 2 + (ny - cy) ** 2 < r ** 2


def _cmap(m: Measurer) -> set[int] | None:
    if m.approximate:
        return None
    return set(TTFont(str(m.font_path)).getBestCmap())


def check_deck(pptx: Path, spec: dict, m: Measurer) -> list[Issue]:
    prs = Presentation(str(pptx))
    cmap = _cmap(m)
    issues: list[Issue] = []
    for n, slide in enumerate(prs.slides, 1):
        layout_name = slide.slide_layout.name
        for shape in slide.shapes:
            name = _logical_name(shape)
            layout_ph = _layout_ph(slide, shape.placeholder_format.idx) if shape.is_placeholder else None
            if layout_name != "01 Title" and _intersects_zone(shape):
                exempt = layout_name == "14 Photo" and (name == "Photo" or name.startswith(ZONE_EXEMPT_ON_PHOTO))
                if not exempt:
                    issues.append(Issue(n, "zone", name, "overlaps the presenter circle"))
            if isinstance(shape, PlaceholderPicture) or shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text
            if not text.strip():
                if shape.is_placeholder:
                    issues.append(Issue(n, "empty", name, "placeholder has no content"))
                continue
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.name not in ALLOWED_FONTS:
                        issues.append(Issue(n, "font", name, f"font '{r.font.name}'"))
            if cmap is not None:
                missing = sorted({ch for ch in text if not ch.isspace() and ord(ch) not in cmap})
                if missing:
                    issues.append(Issue(n, "glyph", name, "missing in Corporate S: " + " ".join(missing)))
            left, top, right, bottom = _insets(shape, layout_ph)
            width = _px(shape.width) - left - right
            height = _px(shape.height) - top - bottom
            measured = m.block_height_px(_paras(shape, layout_ph), width)
            if measured > height + 2:
                issues.append(Issue(n, "overflow", name, f"{measured:.0f}px > {height:.0f}px"))
    return issues
