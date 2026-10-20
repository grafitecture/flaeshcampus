"""Rahmen-Widgets: Leisten-Fortschritt, Evidence, Fußzeile, Fokus-Box, Fotos und Bildbriefings."""
import tempfile
from pathlib import Path

from PIL import Image, ImageOps
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

from . import oxml
from .measure import Measurer
from .tokens import COLOR, EMU_PER_PX, LABEL_SPC, SIZE_PT, px

RAIL_DOT_X = 68          # Mitte der Punkte in der Leiste
RAIL_TOP = 315           # Mitte des ersten Punkts (Punkt-Tops ab y 300)
RAIL_LINE = "3D4B52"


def style(size: str | float, color: str, **extra) -> dict:
    size_pt = SIZE_PT[size] if isinstance(size, str) else size
    return {"size_pt": size_pt, "color": COLOR.get(color, color), **extra}


def textbox(slide, name: str, box, paragraphs: list[tuple[str, dict]], anchor: str = "t"):
    x, y, w, h = box
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tb.name = name
    oxml.fill_text_frame(tb.text_frame, paragraphs, anchor=anchor)
    return tb


def shape(slide, kind, name: str, box, fill: str | None = None, line=None, glow=None, radius_px: float = 0):
    x, y, w, h = box
    sp = slide.shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    sp.name = name
    style_el = sp._element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}style")
    if style_el is not None:
        sp._element.remove(style_el)
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(COLOR.get(fill, fill))
    else:
        sp.fill.background()
    if line:
        color, width, *dash = line
        sp.line.color.rgb = RGBColor.from_string(COLOR.get(color, color))
        sp.line.width = px(width)
        if dash and dash[0]:
            sp.line.dash_style = MSO_LINE.DASH
    else:
        sp.line.fill.background()
    if glow:
        color, radius, alpha = glow
        oxml.add_glow(sp, COLOR.get(color, color), radius, alpha)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = (radius_px or 18) / min(w, h)
    return sp


def rail_progress(slide, chapters: list[str], current: int | None, dark: bool) -> None:
    n = len(chapters)
    pitch = min(104, 600 / (n - 1)) if n > 1 else 0
    centers = [RAIL_TOP + i * pitch for i in range(n)]
    shape(slide, MSO_SHAPE.RECTANGLE, "rail-line", (RAIL_DOT_X - 1, centers[0], 2, centers[-1] - centers[0]),
          fill=RAIL_LINE)
    for i, cy in enumerate(centers, 1):
        if current and i == current:
            size, name = 30, f"rail-dot-{i}-active"
            dot = dict(fill="sky", glow=("sky", 12, 75))   # kleines Leuchten: rund in PowerPoint und LibreOffice
            num_color = "sky"
        elif current and i < current:
            size, name, dot, num_color = 16, f"rail-dot-{i}", dict(fill="steel"), "steel"
        else:
            size, name = 16, f"rail-dot-{i}"
            dot = dict(fill="rail" if dark else "night", line=("dim", 3))
            num_color = "dim"
        shape(slide, MSO_SHAPE.OVAL, name, (RAIL_DOT_X - size / 2, cy - size / 2, size, size), **dot)
        textbox(slide, f"rail-num-{i}", (100, cy - 16, 40, 32), [(str(i), style("small", num_color))], anchor="ctr")


def evidence(slide, text: str, dark: bool) -> None:
    textbox(slide, "evidence-label", (240, 920, 600, 24),
            [("Evidence", style("label", "sky", caps=True, spc=LABEL_SPC))])
    textbox(slide, "evidence-text", (240, 948, 1240, 36),
            [(text, style("small", "muted_dark" if dark else "muted"))])


def footer(slide, text: str, dark: bool) -> None:
    textbox(slide, "footer", (240, 1010, 1200, 30),
            [(text, style("label", "muted_dark" if dark else "steel", caps=True, spc=LABEL_SPC))])


def focus_box(slide, box: tuple[int, int, int], label: str, title: str, text: str, m: Measurer) -> int:
    x, y, w = box
    pad_x, pad_y, inner = 46, 44, w - 92
    label_h = m.line_height_px(SIZE_PT["label"])
    title_h = len(m.wrap(title, SIZE_PT["focus_title"], inner)) * m.line_height_px(SIZE_PT["focus_title"])
    text_h = len(m.wrap(text, SIZE_PT["body"], inner)) * m.line_height_px(SIZE_PT["body"]) if text else 0
    height = round(pad_y + label_h + 10 + title_h + (18 + text_h if text else 0) + pad_y)
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "focus-card", (x, y, w, height),
          fill="night", line=("sky", 2), glow=("sky", 50, 40))
    top = y + pad_y
    textbox(slide, "focus-label", (x + pad_x, top, inner, label_h),
            [(label, style("label", "sky", caps=True, spc=LABEL_SPC))])
    top += label_h + 10
    textbox(slide, "focus-title", (x + pad_x, top, inner, title_h), [(title, style("focus_title", "white"))])
    if text:
        textbox(slide, "focus-text", (x + pad_x, top + title_h + 18, inner, text_h),
                [(text, style("body", "on_dark"))])
    return height


def _as_px(ph) -> tuple[float, float, float, float]:
    return tuple(v / EMU_PER_PX for v in (ph.left, ph.top, ph.width, ph.height))


MIN_PHOTO_RATIO = 0.5  # unter halber Sollgröße gilt ein Foto als zu klein (Spec §6: Platzhalter + Fotoliste)
MAX_PHOTO_RATIO = 2.0  # mehr als doppelte Auflösung bringt nichts und bläht das Deck auf


def _ratio(size: tuple[int, int], box: tuple[float, float, float, float], fit: bool) -> float:
    """Bildpixel je Folienpixel: beim Füllen (beschnitten) zählt die knappere Seite, beim Einpassen die größere."""
    rx, ry = size[0] / box[2], size[1] / box[3]
    return max(rx, ry) if fit else min(rx, ry)


def _usable_image(path: Path, box, fit: bool = False) -> tuple[Path, tuple[int, int]]:
    """Bild am Inhalt erkennen, EXIF-Drehung anwenden, auf höchstens doppelte Rahmenauflösung verkleinern;
    alles außer unverändertem JPEG/PNG wird neu gespeichert. Gibt Pfad und Pixelmaß zurück."""
    with Image.open(path) as im:
        im.load()
        rotated = im.getexif().get(0x0112, 1) != 1
        img = ImageOps.exif_transpose(im) if rotated else im
        scale = MAX_PHOTO_RATIO / _ratio(img.size, box, fit)
        if im.format in ("JPEG", "PNG") and not rotated and scale >= 1:
            return path, im.size
        if scale < 1:
            img = img.resize((max(1, round(img.width * scale)), max(1, round(img.height * scale))), Image.LANCZOS)
        alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
        out = Path(tempfile.mkdtemp(prefix="campus-img-")) / (path.stem + (".png" if alpha else ".jpg"))
        if alpha:
            img.convert("RGBA").save(out, "PNG", optimize=True)
        else:
            img.convert("RGB").save(out, "JPEG", quality=88, optimize=True)
        return out, img.size


def _prepare(photo: dict | None, base_dir: Path, box, warnings: list[str]):
    """Foto prüfen: gibt (Pfad, Pixelmaß) zurück oder den Briefing-Text für eine Platzhalterkarte."""
    photo = photo or {}
    file, brief = photo.get("file"), photo.get("brief")
    path = Path(file) if file else None
    if path is not None and not path.is_absolute():
        path = Path(base_dir) / path
    if path is None or not path.exists():
        if file:
            warnings.append(f"Foto fehlt: {file}")
        return brief or f"Photo missing: {file}"
    fit = bool(photo.get("fit"))
    try:
        usable, size = _usable_image(path, box, fit)
    except Exception as err:  # beschädigt, abgeschnitten, unbekanntes Format: Briefing statt Absturz
        warnings.append(f"Foto unlesbar: {file} ({type(err).__name__})")
        return brief or f"Photo unreadable: {file}"
    if _ratio(size, box, fit) < MIN_PHOTO_RATIO:
        warnings.append(f"Foto zu klein: {file} ({size[0]}×{size[1]} px für {box[2]:.0f}×{box[3]:.0f} px)")
        return brief or f"Photo too small: {file}"
    return usable, size


def _fitted(size: tuple[int, int], box) -> tuple[float, float, float, float]:
    x, y, w, h = box
    scale = min(w / size[0], h / size[1])
    fw, fh = size[0] * scale, size[1] * scale
    return x + (w - fw) / 2, y + (h - fh) / 2, fw, fh


def place_photo(slide, idx: int, photo: dict | None, base_dir: Path, warnings: list[str],
                ellipse: bool = False) -> str | None:
    ph = slide.placeholders[idx]
    box = _as_px(ph)
    prepared = _prepare(photo, base_dir, box, warnings)
    if isinstance(prepared, str):
        return _brief_in_place(slide, ph, prepared)
    usable, size = prepared
    pic = ph.insert_picture(str(usable))
    # Die graue Leerfläche des Platzhalters nicht erben: transparente Bilder zeigen den Folienhintergrund.
    sp_pr = pic._element.spPr
    sp_pr.insert(1 + sum(1 for c in sp_pr if c.tag in (qn("a:xfrm"), qn("a:prstGeom"), qn("a:custGeom"))),
                 OxmlElement("a:noFill"))
    if (photo or {}).get("fit"):  # ganz zeigen: ohne Beschnitt, mittig im Rahmen
        pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
        x, y, w, h = _fitted(size, box)
        pic.left, pic.top, pic.width, pic.height = px(x), px(y), px(w), px(h)
    if ellipse:
        pic.auto_shape_type = MSO_SHAPE.OVAL
        pic.line.color.rgb = RGBColor.from_string(COLOR["sky"])
        pic.line.width = px(6)
        oxml.add_glow(pic, COLOR["sky"], 90, 60)
    return None


def add_photo(slide, name: str, box, photo: dict | None, base_dir: Path, warnings: list[str]) -> str | None:
    """Foto als freies Bild in einen Rahmen setzen (gefüllt und mittig beschnitten oder eingepasst)."""
    prepared = _prepare(photo, base_dir, box, warnings)
    if isinstance(prepared, str):
        brief_card(slide, box, prepared)
        return prepared
    usable, size = prepared
    fit = bool((photo or {}).get("fit"))
    x, y, w, h = _fitted(size, box) if fit else box
    pic = slide.shapes.add_picture(str(usable), px(x), px(y), px(w), px(h))
    pic.name = name
    if not fit:
        image_aspect, box_aspect = size[0] / size[1], w / h
        if image_aspect > box_aspect:
            pic.crop_left = pic.crop_right = (1 - box_aspect / image_aspect) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - image_aspect / box_aspect) / 2
    return None


def _brief_in_place(slide, ph, text: str) -> str:
    box = _as_px(ph)
    tree = ph._element.getparent()
    anchor = list(tree).index(ph._element)
    tree.remove(ph._element)
    count = len(tree)
    brief_card(slide, box, text)
    # Die Karte nimmt die Ebene des Platzhalters ein, damit Texte darüber sichtbar bleiben.
    for offset, el in enumerate(list(tree)[count:]):
        tree.remove(el)
        tree.insert(anchor + offset, el)
    return text


def brief_card(slide, box, brief: str) -> None:
    x, y, w, h = box
    shape(slide, MSO_SHAPE.RECTANGLE, "brief-card", box, fill="mist", line=("steel", 2, True))
    cx, cy = x + w / 2, y + h / 2
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "brief-icon", (cx - 32, cy - 78, 64, 46),
          fill="steel", radius_px=8)
    shape(slide, MSO_SHAPE.OVAL, "brief-icon-lens", (cx - 12, cy - 67, 24, 24), fill="mist")
    textbox(slide, "brief-text", (x + 40, cy - 16, w - 80, min(h / 2, 120)),
            [(brief, style("body", "slate", align="center"))])

