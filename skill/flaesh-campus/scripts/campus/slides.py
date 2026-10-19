"""Ein Renderer pro Folientyp: Layout wählen, Platzhalter füllen, dynamische Elemente zeichnen."""
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from pptx.enum.shapes import PP_PLACEHOLDER

from .layouts import LAYOUTS
from .measure import Measurer
from .plan import default_kicker
from .tokens import PH_IDX, px
from .widgets_content import (bars_card, case_findings, checklist, closing_steps, compare_columns, markers,
                              points_text, quiz_options, scale_caption, scale_chips, steps_grid, summary_items,
                              timeline)
from .widgets_frame import evidence, focus_box, footer, place_photo, rail_progress, zone_ring_on_top

LAYOUT_BY_KEY = {ldef.key: ldef.name for ldef in LAYOUTS}
DARK_KEYS = {"cover", "chapter", "hook", "closing", "photo", "quiz"}
NO_FOOTER = {"cover", "chapter", "photo"}


@dataclass
class BuildReport:
    slides: int = 0
    briefs: list[tuple[int, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class Ctx:
    prs: object
    plan: dict
    m: Measurer
    base_dir: Path
    report: BuildReport


# ---------------------------------------------------------------- Grundfunktionen

def new_slide(ctx: Ctx, layout_name: str):
    layout = next(lay for lay in ctx.prs.slide_layouts if lay.name == layout_name)
    slide = ctx.prs.slides.add_slide(layout)
    ctx.report.slides += 1
    return slide


def _placeholder(slide, name: str):
    for ph in slide.placeholders:
        if name == "Title":
            if ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                return ph
        elif ph.placeholder_format.idx == PH_IDX[name]:
            return ph
    return None


def fill(slide, name: str, text: str | None) -> None:
    ph = _placeholder(slide, name)
    if ph is None:
        return
    if not text or not str(text).strip():
        ph._element.getparent().remove(ph._element)
        return
    tf = ph.text_frame
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    lines = str(text).split("\n")
    tf.paragraphs[0].add_run().text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().add_run().text = line


def _photo(ctx: Ctx, slide, name: str, photo: dict | None, ellipse: bool = False) -> None:
    brief = place_photo(slide, PH_IDX[name], photo, ctx.base_dir, ctx.report.warnings, ellipse=ellipse)
    if brief:
        ctx.report.briefs.append((ctx.report.slides, brief))


def _frame(ctx: Ctx, slide, s: dict, chapter: int | None = None) -> None:
    ep = ctx.plan["episode"]
    key = s["type"]
    dark = key in DARK_KEYS
    fill(slide, "Episode", f"{ep['number']:02d}")
    if key != "cover":
        rail_progress(slide, ep["chapters"], chapter if chapter is not None else s.get("chapter"), dark)
    if s.get("evidence"):
        evidence(slide, s["evidence"], dark)
    if key not in NO_FOOTER:
        footer(slide, f"{ep['title'].upper()} · {ctx.report.slides}", dark)
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"]


def _start(ctx: Ctx, s: dict, title: bool = True):
    slide = new_slide(ctx, LAYOUT_BY_KEY[s["type"]])
    fill(slide, "Kicker", s.get("kicker") or default_kicker(ctx.plan, s))
    if title:
        fill(slide, "Title", s.get("title"))
    return slide


def _points(slide, points: list[dict] | None) -> None:
    ph = _placeholder(slide, "Points")
    if ph is None:
        return
    if points:
        points_text(ph, points)
    else:
        ph._element.getparent().remove(ph._element)


def _focus(ctx: Ctx, slide, s: dict) -> None:
    focus = s["focus"]
    focus_box(slide, (1220, 330, 580), focus.get("label", "Clinical focus"), focus["title"],
              focus.get("text", ""), ctx.m)


# ---------------------------------------------------------------- Renderer

def render_cover(ctx: Ctx, s: dict) -> None:
    ep = ctx.plan["episode"]
    presenter = ep.get("presenter") or {}
    slide = new_slide(ctx, LAYOUT_BY_KEY["cover"])
    fill(slide, "Kicker", s.get("kicker") or f"Online Training · Episode {ep['number']:02d}")
    fill(slide, "Title", ep["title"])
    fill(slide, "Subtitle", ep.get("subtitle"))
    fill(slide, "Presenter", presenter.get("name"))
    fill(slide, "Credentials", presenter.get("credentials"))
    _photo(ctx, slide, "Photo", presenter.get("photo") or {"brief": f"Portrait {presenter.get('name', '')}".strip()},
           ellipse=True)
    _frame(ctx, slide, s)


def render_chapter(ctx: Ctx, s: dict) -> None:
    chapters = ctx.plan["episode"]["chapters"]
    n = s["chapter"]
    slide = new_slide(ctx, LAYOUT_BY_KEY["chapter"])
    fill(slide, "Number", str(n))
    fill(slide, "Kicker", s.get("kicker") or f"Phase {n} of {len(chapters)}")
    fill(slide, "Title", s.get("title") or chapters[n - 1])
    fill(slide, "Line", s.get("line"))
    _frame(ctx, slide, s, chapter=n)


def render_hook(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s, title=False)
    quote = s["quote"].strip()
    if not quote.startswith(("“", '"')):
        quote = f"“{quote}”"
    fill(slide, "Quote", quote)
    fill(slide, "TruthLabel", s["truth"].get("label", "The hard truth"))
    fill(slide, "TruthText", s["truth"]["text"])
    fill(slide, "AnswerLabel", s["answer"].get("label", "The fläsh answer"))
    fill(slide, "AnswerText", s["answer"]["text"])
    _frame(ctx, slide, s)


def render_closing(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    closing_steps(slide, s["steps"])
    nxt = s["next"]
    fill(slide, "NextLabel", nxt.get("label", "Up next"))
    fill(slide, "NextTitle", nxt["title"])
    fill(slide, "NextText", nxt.get("text"))
    _frame(ctx, slide, s)


def render_points(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _points(slide, s["points"])
    if s.get("focus"):
        ph = _placeholder(slide, "Points")
        left, top, height = ph.left, ph.top, ph.height
        ph.left, ph.top, ph.width, ph.height = left, top, px(900), height
        _focus(ctx, slide, s)
    _frame(ctx, slide, s)


def render_points_photo(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _points(slide, s["points"])
    _photo(ctx, slide, "Photo", s["photo"])
    fill(slide, "Caption", s.get("caption"))
    _frame(ctx, slide, s)


def render_points_chart(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _points(slide, s["points"])
    bars_card(slide, s["chart"])
    _frame(ctx, slide, s)


def render_statement(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s, title=False)
    fill(slide, "Big", s.get("big"))
    fill(slide, "Statement", s.get("statement"))
    fill(slide, "Text", s.get("text"))
    _frame(ctx, slide, s)


def render_checklist(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    checklist(slide, s["items"], w=900 if s.get("focus") else 1260)
    if s.get("focus"):
        _focus(ctx, slide, s)
    _frame(ctx, slide, s)


def render_steps(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    steps_grid(slide, s["steps"], dark=False)
    _frame(ctx, slide, s)


def render_compare(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    compare_columns(slide, s["left"], s["right"], s.get("style", "do_dont"))
    _frame(ctx, slide, s)


def render_scale(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    if "scale" in s:
        scale_chips(slide, s["scale"]["items"], s["scale"].get("highlight", []))
    else:
        timeline(slide, s["timeline"]["points"])
    if s.get("caption"):
        scale_caption(slide, s["caption"])
    _points(slide, s.get("points"))
    _frame(ctx, slide, s)


def render_summary(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    summary_items(slide, s["items"])
    _frame(ctx, slide, s)


def render_photo(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _photo(ctx, slide, "Photo", s["photo"])
    fill(slide, "Caption", s.get("caption"))
    zone_ring_on_top(slide, dark=True)
    _frame(ctx, slide, s)


def render_before_after(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _photo(ctx, slide, "PhotoBefore", s["before"]["photo"])
    _photo(ctx, slide, "PhotoAfter", s["after"]["photo"])
    fill(slide, "LabelBefore", s["before"]["label"])
    fill(slide, "LabelAfter", s["after"]["label"])
    _frame(ctx, slide, s)


def render_photo_annotated(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _photo(ctx, slide, "Photo", s["photo"])
    markers(slide, (240, 300, 900, 600), s["markers"])
    _frame(ctx, slide, s)


def render_case(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    fill(slide, "Story", s["story"])
    case_findings(slide, s["findings"])
    _frame(ctx, slide, s)


def render_quiz(ctx: Ctx, s: dict) -> None:
    for reveal in (False, True):
        slide = new_slide(ctx, LAYOUT_BY_KEY["quiz"])
        fill(slide, "Kicker", s.get("kicker") or ("Knowledge check · Answer" if reveal else "Knowledge check"))
        fill(slide, "Question", s["question"])
        quiz_options(slide, s["options"], s["answer"] if reveal else None, s.get("explanation") if reveal else None)
        notes = dict(s)
        if reveal:
            notes["notes"] = s.get("explanation") or s.get("notes")
        _frame(ctx, slide, notes)


RENDERERS: dict[str, Callable[[Ctx, dict], None]] = {
    "cover": render_cover, "chapter": render_chapter, "hook": render_hook, "closing": render_closing,
    "points": render_points, "points_photo": render_points_photo, "points_chart": render_points_chart,
    "statement": render_statement, "checklist": render_checklist, "steps": render_steps,
    "compare": render_compare, "scale": render_scale, "summary": render_summary, "photo": render_photo,
    "before_after": render_before_after, "photo_annotated": render_photo_annotated, "case": render_case,
    "quiz": render_quiz,
}
