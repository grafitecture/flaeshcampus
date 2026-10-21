"""Erzeugt die fläsh-Campus-Vorlage: Theme, Master und 18 Layouts."""
import json
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.slide import SlideLayoutPart
from pptx.shapes.shapetree import SlideShapes

from . import oxml
from .layouts import LAYOUTS, DecorDef, LayoutDef, PhDef, frame_decor
from .tokens import COLOR, FONT, PH_IDX, SLIDE_EMU, px

THEME_COLORS = {
    "dk1": "333333", "lt1": "FFFFFF", "dk2": "1F272B", "lt2": "DEE7EC",
    "accent1": "00A5FF", "accent2": "5D6C73", "accent3": "ABB7BE", "accent4": "283237",
    "accent5": "7D8B91", "accent6": "C9D4DA", "hlink": "00A5FF", "folHlink": "5D6C73",
}
FIRST_LAYOUT_ID = 2147483649

_NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')

LAYOUT_XML = (
    f'<p:sldLayout {_NS} preserve="1" userDrawn="1"><p:cSld name="{{name}}"><p:spTree>'
    '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
    '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/>'
    '<a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:spTree></p:cSld>'
    '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>'
)


def _rpr(size: int, color: str, font: str) -> str:
    return (f'<a:defRPr sz="{size}" kern="1200"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="+{font}-lt"/><a:ea typeface="+{font}-ea"/><a:cs typeface="+{font}-cs"/></a:defRPr>')


def _ppr(level: int, size: int, color: str, font: str) -> str:
    return (f'<a:lvl{level}pPr marL="0" indent="0" algn="l" defTabSz="914400" rtl="0" eaLnBrk="1" '
            f'latinLnBrk="0" hangingPunct="1"><a:lnSpc><a:spcPct val="100000"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>{_rpr(size, color, font)}</a:lvl{level}pPr>')


TXSTYLES_XML = (
    f'<p:txStyles {_NS}>'
    f'<p:titleStyle>{_ppr(1, 3900, COLOR["slate"], "mj")}</p:titleStyle>'
    f'<p:bodyStyle>{_ppr(1, 2100, COLOR["slate"], "mn")}{_ppr(2, 1700, COLOR["ink"], "mn")}'
    f'{_ppr(3, 1700, COLOR["ink"], "mn")}</p:bodyStyle>'
    f'<p:otherStyle><a:defPPr><a:defRPr lang="en-US"/></a:defPPr>{_ppr(1, 1700, COLOR["ink"], "mn")}</p:otherStyle>'
    '</p:txStyles>'
)


def build_template(out: Path, assets_dir: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_EMU
    master = prs.slide_master
    _apply_theme(master)
    _style_master(master)
    _remove_layouts(master)
    for i, ldef in enumerate(LAYOUTS):
        layout = _new_layout(prs, master, ldef.name, FIRST_LAYOUT_ID + i)
        _populate(layout, ldef, Path(assets_dir))
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def _apply_theme(master) -> None:
    part = master.part.part_related_by(RT.THEME)
    xml = etree.fromstring(part.blob)
    xml.set("name", "fläsh Campus")
    scheme = xml.find(".//" + qn("a:clrScheme"))
    scheme.set("name", "fläsh Campus")
    for key, value in THEME_COLORS.items():
        slot = scheme.find(qn(f"a:{key}"))
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, qn("a:srgbClr")).set("val", value)
    fonts = xml.find(".//" + qn("a:fontScheme"))
    fonts.set("name", "fläsh Campus")
    for tag in ("a:majorFont", "a:minorFont"):
        fonts.find(qn(tag)).find(qn("a:latin")).set("typeface", FONT)
    part._blob = etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)


def _style_master(master) -> None:
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = RGBColor.from_string(COLOR["white"])
    old = master._element.find(qn("p:txStyles"))
    new = parse_xml(TXSTYLES_XML)
    old.addprevious(new)
    master._element.remove(old)


def _remove_layouts(master) -> None:
    lst = master._element.get_or_add_sldLayoutIdLst()
    for el in list(lst):
        rId = el.get(qn("r:id"))
        lst.remove(el)
        master.part.drop_rel(rId)


def _new_layout(prs, master, name: str, layout_id: int):
    package = prs.part.package
    partname = package.next_partname("/ppt/slideLayouts/slideLayout%d.xml")
    element = parse_xml(LAYOUT_XML.format(name=escape(name, {'"': "&quot;"})))
    part = SlideLayoutPart(partname, CT.PML_SLIDE_LAYOUT, package, element)
    part.relate_to(master.part, RT.SLIDE_MASTER)
    rId = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    entry = OxmlElement("p:sldLayoutId")
    entry.set("id", str(layout_id))
    entry.set(qn("r:id"), rId)
    master._element.get_or_add_sldLayoutIdLst().append(entry)
    return part.slide_layout


def _populate(layout, ldef: LayoutDef, assets_dir: Path) -> None:
    layout.background.fill.solid()
    bg = COLOR["white"] if ldef.frame == "light" else COLOR["night"]
    layout.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    shapes = SlideShapes(layout.shapes._spTree, layout)
    behind = [d for d in ldef.decor if d.style.get("behind")]
    front = [d for d in ldef.decor if not d.style.get("behind")]
    for d in behind + frame_decor(ldef.frame) + front:
        _add_decor(shapes, d, assets_dir)
    for ph in ldef.placeholders:
        _add_placeholder(shapes, ph)


def _box(box) -> tuple[int, int, int, int]:
    return tuple(px(v) for v in box)


def _strip_style(shape) -> None:
    style = shape._element.find(qn("p:style"))
    if style is not None:
        shape._element.remove(style)


def _apply_fill_line(shape, fill: str | None, fill_alpha, line, glow) -> None:
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if fill_alpha is not None:
            oxml.set_fill_alpha(shape, fill_alpha)
    else:
        shape.fill.background()
    if line:
        color, width, *rest = line
        shape.line.color.rgb = RGBColor.from_string(color)
        shape.line.width = px(width)
        if rest and rest[0] is not None:
            oxml.set_line_alpha(shape, rest[0])
    else:
        shape.line.fill.background()
    if glow:
        oxml.add_glow(shape, *glow)


def _add_decor(shapes, d: DecorDef, assets_dir: Path) -> None:
    x, y, w, h = _box(d.box)
    st = d.style
    if d.kind == "image":
        pic = shapes.add_picture(str(assets_dir / st["file"]), x, y, w, h)
        pic.name = d.name
        if st.get("ellipse"):
            pic.auto_shape_type = MSO_SHAPE.OVAL
        return
    if d.kind == "text":
        tb = shapes.add_textbox(x, y, w, h)
        tb.name = d.name
        oxml.fill_text_frame(tb.text_frame, [(st["text"], st["level"])], anchor=st.get("anchor", "t"))
        return
    mso = {"rect": MSO_SHAPE.RECTANGLE, "roundrect": MSO_SHAPE.ROUNDED_RECTANGLE,
           "ellipse": MSO_SHAPE.OVAL, "radial": MSO_SHAPE.OVAL}[d.kind]
    shape = shapes.add_shape(mso, x, y, w, h)
    shape.name = d.name
    _strip_style(shape)
    if d.kind == "radial":
        oxml.set_radial_fill(shape, st["color"], st["stops"])
        shape.line.fill.background()
    else:
        _apply_fill_line(shape, st.get("fill"), st.get("fill_alpha"), st.get("line"), st.get("glow"))
    if d.kind == "roundrect":
        shape.adjustments[0] = st.get("radius_px", 18) / min(d.box[2], d.box[3])


def _add_placeholder(shapes, ph: PhDef) -> None:
    x, y, w, h = _box(ph.box)
    if ph.ph_type == "pic":
        shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    else:
        shape = shapes.add_textbox(x, y, w, h)
    idx = None if ph.ph_type == "title" else PH_IDX[ph.name]
    oxml.to_placeholder(shape, ph.ph_type, idx, ph.name, ph.levels, ph.prompt,
                        ph.anchor, ph.insets, ph.geom)
    fill = ph.fill or (COLOR["mist"] if ph.ph_type == "pic" else None)
    _apply_fill_line(shape, fill, ph.fill_alpha, ph.line, ph.glow)


def layout_spec() -> dict:
    """Maße und Stile aller Layouts als JSON-taugliches Dict (Quelle für Prüfung und Skill-Doku)."""
    spec = {}
    for ldef in LAYOUTS:
        spec[ldef.name] = {
            "key": ldef.key,
            "frame": ldef.frame,
            "zone": ldef.frame != "cover",
            "placeholders": {
                ph.name: {
                    "idx": 0 if ph.ph_type == "title" else PH_IDX[ph.name],
                    "type": ph.ph_type,
                    "box": list(ph.box),
                    "levels": ph.levels,
                    "insets": list(ph.insets),
                }
                for ph in ldef.placeholders
            },
        }
    return json.loads(json.dumps(spec))


def write_layout_spec(path: Path) -> None:
    Path(path).write_text(json.dumps(layout_spec(), indent=2, ensure_ascii=False), encoding="utf-8")
