"""Baut eine fläsh-Campus-PowerPoint aus einem Folienplan."""
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

from .measure import Measurer, find_font
from .plan import PlanError, validate_plan
from .slides import RENDERERS, BuildReport, Ctx


def clear_slides(prs) -> None:
    sld_id_lst = prs.slides._sldIdLst
    for sld_id in list(sld_id_lst):
        rId = sld_id.get(qn("r:id"))
        sld_id_lst.remove(sld_id)
        prs.part.drop_rel(rId)


def build_deck(plan: dict, template: Path, out: Path, base_dir: Path, m: Measurer | None = None) -> BuildReport:
    errors = validate_plan(plan)
    if errors:
        raise PlanError(errors)
    prs = Presentation(str(template))
    clear_slides(prs)
    report = BuildReport()
    ctx = Ctx(prs, plan, m or Measurer(find_font()), Path(base_dir), report)
    for slide in plan["slides"]:
        RENDERERS[slide["type"]](ctx, slide)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return report
