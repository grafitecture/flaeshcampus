"""Inhalts-Widgets: Punkte, Checkliste, Schritte, Balken, Vergleich, Skala, Zeitstrahl,
Zusammenfassung, Marker, Befunde, Quiz, Abschluss-Schritte. Symbole sind Formen, keine Zeichen."""
from typing import Literal

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Pt

from . import oxml
from .tokens import COLOR, KICKER_SPC, LABEL_SPC, px
from .widgets_frame import shape, style, textbox


# ---------------------------------------------------------------- Hilfen

def _number(slide, name: str, box, n: int | str, fill: str, color: str, size: float,
            line=None, glow=None):
    """Kreis mit Zahl/Buchstabe als Formtext."""
    sp = shape(slide, MSO_SHAPE.OVAL, name, box, fill=fill, line=line, glow=glow)
    oxml.fill_text_frame(sp.text_frame, [(str(n), style(size, color, align="center"))], anchor="ctr")
    return sp


def _lines_group(slide, name: str, segments, color: str, width_px: float):
    """Icon aus geraden Linien (in px); als Gruppe benannt."""
    group = slide.shapes.add_group_shape()
    group.name = name
    for (x1, y1), (x2, y2) in segments:
        line = group.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, px(x1), px(y1), px(x2), px(y2))
        line.line.color.rgb = RGBColor.from_string(COLOR.get(color, color))
        line.line.width = px(width_px)
    return group


def check_icon(slide, name: str, x: float, y: float, size: float, color: str = "sky"):
    s = size
    return _lines_group(slide, name, [((x + .15 * s, y + .52 * s), (x + .40 * s, y + .78 * s)),
                                      ((x + .40 * s, y + .78 * s), (x + .86 * s, y + .22 * s))], color, s * .12)


def cross_icon(slide, name: str, x: float, y: float, size: float, color: str = "slate"):
    s = size
    return _lines_group(slide, name, [((x + .2 * s, y + .2 * s), (x + .8 * s, y + .8 * s)),
                                      ((x + .8 * s, y + .2 * s), (x + .2 * s, y + .8 * s))], color, s * .12)


# ---------------------------------------------------------------- Punkte

def points_text(ph, points: list[dict]) -> None:
    tf = ph.text_frame
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    first = tf.paragraphs[0]
    for run in list(first.runs):
        run._r.getparent().remove(run._r)
    for i, point in enumerate(points):
        lead = first if i == 0 else tf.add_paragraph()
        lead.level = 0
        num = lead.add_run()
        num.text = f"{i + 1:02d}"
        num.font.color.rgb = RGBColor.from_string(COLOR["sky"])
        lead.add_run().text = "\t" + point["lead"]
        if point.get("text"):
            body = tf.add_paragraph()
            body.level = 1
            body.add_run().text = point["text"]
    first.space_before = Pt(0)


# ---------------------------------------------------------------- Checkliste, Schritte, Zusammenfassung

def checklist(slide, items: list[str], x: float = 240, y: float = 330, w: float = 900) -> None:
    for i, item in enumerate(items, 1):
        top = y + (i - 1) * 82
        _number(slide, f"check-num-{i}", (x, top, 60, 60), i, "mist", "sky", 16)
        textbox(slide, f"check-text-{i}", (x + 86, top, w - 86, 60), [(item, style("body", "ink"))], anchor="ctr")


def steps_grid(slide, steps: list[dict], dark: bool, area=(240, 330, 1260, 560)) -> None:
    ax, ay, aw, ah = area
    rows = {3: [3], 4: [2, 2], 5: [3, 2]}.get(len(steps), [len(steps)])
    cols = max(rows)
    cw, rh = aw / cols, ah / len(rows)
    title_color, text_color = ("white", "on_dark") if dark else ("slate", "ink")
    i = 0
    for r, count in enumerate(rows):
        for c in range(count):
            step = steps[i]
            i += 1
            x, y = ax + c * cw, ay + r * rh
            _number(slide, f"step-num-{i}", (x, y, 56, 56), i, "sky", "white", 18)
            textbox(slide, f"step-title-{i}", (x + 76, y + 6, cw - 106, 48), [(step["title"], style("lead", title_color))])
            textbox(slide, f"step-text-{i}", (x + 76, y + 60, cw - 106, rh - 80), [(step["text"], style("body", text_color))])


def summary_items(slide, items: list[dict]) -> None:
    for i, item in enumerate(items, 1):
        col, row = (0, i - 1) if i <= 3 else (1, i - 4)
        x, y = 240 + col * 640, 330 + row * 200
        shape(slide, MSO_SHAPE.OVAL, f"sum-check-{i}", (x, y, 48, 48), fill="sky")
        check_icon(slide, f"sum-check-icon-{i}", x + 8, y + 8, 32, "white")
        textbox(slide, f"sum-title-{i}", (x + 72, y, 528, 50), [(item["title"], style("lead", "slate"))])
        textbox(slide, f"sum-text-{i}", (x + 72, y + 54, 528, 120), [(item["text"], style("body", "ink"))])


def closing_steps(slide, steps: list[dict]) -> None:
    for i, step in enumerate(steps, 1):
        y = 330 + (i - 1) * 160
        _number(slide, f"next-step-num-{i}", (240, y, 56, 56), i, "sky", "white", 18)
        textbox(slide, f"next-step-title-{i}", (316, y + 3, 804, 50), [(step["title"], style("lead", "white"))])
        textbox(slide, f"next-step-text-{i}", (316, y + 56, 804, 80), [(step["text"], style("body", "on_dark"))])


# ---------------------------------------------------------------- Balken

def bars_card(slide, chart: dict, box=(1000, 330, 800, 364)) -> None:
    x, y, w, _ = box
    bars = chart["bars"]
    pitch, note_h = 104, (40 if chart.get("note") else 0)
    h = 34 + 34 + 16 + len(bars) * pitch + note_h + 20
    if y + h > 700:                       # sonst kollidiert die Karte mit der Kreis-Zone
        w = 540
    shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, "chart-card", (x, y, w, h), fill="mist", radius_px=18)
    inner = w - 100
    textbox(slide, "chart-title", (x + 50, y + 34, inner, 30),
            [(chart.get("title", ""), style("small", "slate", caps=True, spc=KICKER_SPC))])
    track_w = inner - 130
    tops = [y + 84 + (i - 1) * pitch for i in range(1, len(bars) + 1)]
    for i, (bar, by) in enumerate(zip(bars, tops), 1):
        track_y = by + 46                  # Beschriftung (40 px, eine Zeile 17 pt) steht über dem Balken
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, f"bar-track-{i}", (x + 50, track_y, track_w, 40),
              fill="white", radius_px=20)
        hi = bar.get("highlight")
        fill_w = max(track_w * bar["value"] / 100, 1)
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, f"bar-fill-{i}", (x + 50, track_y, fill_w, 40),
              fill="sky" if hi else "steel", glow=("sky", 30, 60) if hi else None, radius_px=20)
        value = bar.get("display", f"{bar['value']:g}%")
        textbox(slide, f"bar-value-{i}", (x + w - 50 - 120, track_y - 5, 120, 50),
                [(value, style("lead", "slate", align="right"))], anchor="ctr")
    for i, (bar, by) in enumerate(zip(bars, tops), 1):  # Text zuletzt: das Leuchten liegt nie über der Schrift
        textbox(slide, f"bar-label-{i}", (x + 50, by, inner, 40), [(bar["label"], style("body", "slate"))])
    if chart.get("note"):
        textbox(slide, "chart-note", (x + 50, y + h - 20 - 36, inner, 36), [(chart["note"], style("small", "muted"))])


# ---------------------------------------------------------------- Vergleich

def compare_columns(slide, left: dict, right: dict, style_: Literal["do_dont", "myth_fact"] = "do_dont") -> None:
    positive = {"do_dont": "l", "myth_fact": "r"}[style_]
    for side, data, x in (("l", left, 240), ("r", right, 900)):
        good = side == positive
        head = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, f"compare-head-{side}", (x, 330, 600, 64),
                     fill="sky" if good else "slate", radius_px=14)
        oxml.fill_text_frame(head.text_frame, [(data["label"], style("kicker", "white", caps=True, spc=KICKER_SPC))],
                             anchor="ctr")
        head.text_frame.margin_left = px(28)
        for i, item in enumerate(data["items"], 1):
            y = 430 + (i - 1) * 104
            if good:
                check_icon(slide, f"icon-check-{side}-{i}", x, y + 6, 36, "sky")
            else:
                cross_icon(slide, f"icon-cross-{side}-{i}", x, y + 6, 36, "slate")
            textbox(slide, f"compare-item-{side}-{i}", (x + 56, y, 544, 96), [(item, style("body", "ink"))])


# ---------------------------------------------------------------- Skala und Zeitstrahl

def scale_chips(slide, items: list[str], highlight: list[str], box=(240, 330, 1560, 140)) -> None:
    x, y, w, h = box
    gap = 8
    cw = (w - gap * (len(items) - 1)) / len(items)
    for i, item in enumerate(items, 1):
        hi = item in (highlight or [])
        chip = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, f"scale-chip-{i}", (x + (i - 1) * (cw + gap), y, cw, h - 40),
                     fill="sky" if hi else "mist", glow=("sky", 26, 60) if hi else None, radius_px=10)
        oxml.fill_text_frame(chip.text_frame, [(item, style("body", "white" if hi else "slate", align="center"))],
                             anchor="ctr")


def scale_caption(slide, text: str, y: float = 490) -> None:
    textbox(slide, "scale-caption", (240, y, 1260, 40), [(text, style("small", "muted"))])


def timeline(slide, points: list[dict], box=(240, 330, 1260, 200), dark: bool = False) -> None:
    x, y, w, _ = box
    inner_x, inner_w = x + 150, w - 300
    n = len(points)
    xs = [inner_x + (inner_w * i / (n - 1) if n > 1 else inner_w / 2) for i in range(n)]
    line_y = y + 60
    label_w = min(300, inner_w / (n - 1) - 24) if n > 1 else 300  # Nachbarn überlappen nicht (bis 6 Punkte)
    shape(slide, MSO_SHAPE.RECTANGLE, "tl-line", (xs[0], line_y - 2, xs[-1] - xs[0], 4), fill="steel")
    for i, (cx, point) in enumerate(zip(xs, points), 1):
        shape(slide, MSO_SHAPE.OVAL, f"tl-dot-{i}", (cx - 12, line_y - 12, 24, 24), fill="sky", glow=("sky", 10, 60))
        textbox(slide, f"tl-at-{i}", (cx - label_w / 2, y - 10, label_w, 50),
                [(point["at"], style("lead", "white" if dark else "slate", align="center"))], anchor="b")
        textbox(slide, f"tl-label-{i}", (cx - label_w / 2, line_y + 26, label_w, 80),
                [(point["label"], style("body", "on_dark" if dark else "ink", align="center"))])


# ---------------------------------------------------------------- Fotomarker, Befunde, Quiz

def markers(slide, photo_box, items: list[dict]) -> None:
    px0, py0, pw, ph = photo_box
    for i, mk in enumerate(items, 1):
        cx, cy = px0 + mk["x"] * pw, py0 + mk["y"] * ph
        _number(slide, f"marker-{i}", (cx - 24, cy - 24, 48, 48), i, "sky", "white", 18,
                line=("white", 3), glow=("sky", 20, 70))
        y = 320 + (i - 1) * 76
        _number(slide, f"legend-num-{i}", (1200, y + 8, 44, 44), i, "sky", "white", 16)
        textbox(slide, f"legend-{i}", (1264, y, 496, 60), [(mk["label"], style("body", "ink"))], anchor="ctr")


def case_findings(slide, findings: list[dict], y: float = 560) -> None:
    textbox(slide, "finding-head", (240, y - 44, 560, 30),
            [("Finding", style("label", "sky", caps=True, spc=LABEL_SPC))])
    textbox(slide, "cause-head", (880, y - 44, 620, 30),
            [("Cause", style("label", "sky", caps=True, spc=LABEL_SPC))])
    for i, item in enumerate(findings, 1):
        top = y + (i - 1) * 72
        textbox(slide, f"finding-{i}", (240, top, 560, 64), [(item["finding"], style("body", "slate"))], anchor="ctr")
        shape(slide, MSO_SHAPE.RIGHT_ARROW, f"arrow-{i}", (820, top + 20, 40, 24), fill="sky")
        textbox(slide, f"cause-{i}", (880, top, 620, 64), [(item["cause"], style("body", "ink"))], anchor="ctr")


def quiz_options(slide, options: list[str], answer: int | None, explanation: str | None) -> None:
    for i, option in enumerate(options, 1):
        y = 480 + (i - 1) * 120
        correct = answer is not None and i - 1 == answer
        shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, f"quiz-opt-{i}", (240, y, 1260, 100),
              fill="sky" if correct else "night2", glow=("sky", 50, 50) if correct else None, radius_px=18)
        _number(slide, f"quiz-letter-{i}", (270, y + 20, 60, 60), "ABC"[i - 1],
                "white" if correct else "night2", "night" if correct else "sky", 18,
                line=None if correct else ("sky", 3))
        textbox(slide, f"quiz-text-{i}", (360, y, 1100, 100),
                [(option, style("lead", "night" if correct else "white"))], anchor="ctr")
    if answer is not None and explanation:
        textbox(slide, "quiz-expl", (240, 880, 1260, 80), [(explanation, style("body", "on_dark"))])
